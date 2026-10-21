import unittest
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import Depends, FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api.deps import get_current_user, require_permission
from app.api.routers import sales_profit as router_module
from app.db.ods import get_ods_db
from app.db.sales_profit import get_sales_profit_db
from app.db.session import get_db
from app.models.ads import AdsPublishBatch
from app.services import sales_profit as service


def source_row(customer="A", quantity="10", amount="100", sale_quantity=None, sale_amount=None,
               product="SKU-1", channel="线下 A", product_type="正装", brand="品牌 A", **extra):
    return {
        "customer_code": customer, "customer_name": f"客户 {customer}" if customer else "未识别客户",
        "product_code": product, "product_name": f"商品 {product}", "product_type": product_type,
        "brand": brand, "channel": channel, "quantity": quantity, "paid_amount": amount,
        "sale_quantity": sale_quantity if sale_quantity is not None else quantity,
        "sale_amount": sale_amount if sale_amount is not None else amount,
        "return_quantity": "0", "return_amount": "0", **extra,
    }


DIMENSIONS = {
    "线下 A": {"owner": "销售 A", "channel_type": "线下"},
    "线下 B": {"owner": "销售 B", "channel_type": "线下"},
    "线下 A2": {"owner": "销售 A", "channel_type": "线下"},
    "线上 C": {"owner": "销售 A", "channel_type": "线上"},
}


class SalesProfitCalculationTest(unittest.TestCase):
    def enriched(self, raw, prices=None, params=None, previous=None):
        params = params or service.AnalysisParams(view="prices")
        rows = service.prepare_rows(raw, DIMENSIONS, params)
        prior = service.prepare_rows(previous or [], DIMENSIONS, params)
        return service.enrich_rows(rows, prices or {}, prior, params)

    def test_weighted_company_price_uses_other_owners_and_positive_sales(self):
        params = service.AnalysisParams(view="prices", owner="销售 A")
        rows = self.enriched([
            source_row(amount="80"), source_row(customer="B", quantity="30", amount="360", channel="线下 B"),
        ], {"SKU-1": Decimal("6")}, params)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["reference_price"], Decimal("11"))
        self.assertEqual(rows[0]["price_delta"], Decimal("-3"))
        self.assertEqual(rows[0]["price_effect_amount"], Decimal("-30"))
        self.assertAlmostEqual(float(rows[0]["price_delta_rate"]), -27.27272727)

    def test_returns_are_excluded_from_profit_and_price(self):
        rows = self.enriched([source_row(
            quantity="8", amount="76", sale_quantity="10", sale_amount="100",
            return_quantity="-2", return_amount="-24",
        )], {"SKU-1": Decimal("6")})
        self.assertEqual(rows[0]["avg_price"], Decimal("10"))
        self.assertEqual(rows[0]["estimated_cost"], Decimal("60"))
        self.assertEqual(rows[0]["estimated_profit"], Decimal("40"))
        self.assertAlmostEqual(service.summarize(rows)["margin_rate"], 40)

    def test_screenshot_zero_amount_returns_do_not_inflate_profit(self):
        rows = self.enriched([
            source_row(product='LE847701', quantity='-1694', amount='5565.10', sale_quantity='109', sale_amount='5565.10', return_quantity='1803'),
            source_row(product='LF148800', quantity='-493', amount='548', sale_quantity='5', sale_amount='548', return_quantity='498'),
            source_row(product='RETURN', quantity='-50', amount='0', sale_quantity='0', sale_amount='0', return_quantity='50'),
        ], {'LE847701': '45', 'LF148800': '80', 'RETURN': '20'}, service.AnalysisParams(view='profit'))
        self.assertEqual(rows[0]['estimated_profit'], Decimal('660.10'))
        self.assertEqual(rows[1]['estimated_profit'], Decimal('148'))
        self.assertAlmostEqual(float(rows[0]['margin_rate']), 11.86142208)
        self.assertIsNone(rows[2]['estimated_profit'])
        self.assertIsNone(rows[2]['avg_price'])
        self.assertEqual(rows[0]['quantity'], Decimal('-1694'))
        summary = service.summarize(rows)
        self.assertEqual(summary['estimated_profit'], 808.10)
        self.assertEqual(summary['estimated_cost'], 5305)
        self.assertEqual(service.breakdown(rows, 'owner')[0]['estimated_profit'], 808.10)

    def test_missing_cost_never_becomes_zero_or_changes_margin_denominator(self):
        rows = self.enriched([source_row(), source_row(product="SKU-2", amount="300")], {"SKU-1": "6"})
        self.assertIsNone(rows[1]["estimated_profit"])
        summary = service.summarize(rows)
        self.assertEqual(summary["paid_amount"], 400)
        self.assertEqual(summary["priced_amount"], 100)
        self.assertEqual(summary["missing_cost_amount"], 300)
        self.assertEqual(summary["estimated_profit"], 40)
        self.assertEqual(summary["margin_rate"], 40)
        self.assertEqual(summary["cost_coverage_rate"], 25)

    def test_margin_does_not_divide_nonpositive_revenue(self):
        rows = self.enriched([source_row(quantity="-2", amount="-20", sale_quantity="0", sale_amount="0")], {"SKU-1": "6"})
        self.assertIsNone(rows[0]["margin_rate"])
        self.assertIsNone(rows[0]["avg_price"])
        self.assertIsNone(rows[0]["estimated_cost"])
        self.assertIsNone(rows[0]["estimated_profit"])
        self.assertIsNone(service.summarize(rows)["margin_rate"])
        self.assertIsNone(service.summarize(rows)["cost_coverage_rate"])

    def test_anonymous_customer_and_missing_sku_keep_revenue_but_no_comparison(self):
        rows = self.enriched([source_row(customer=""), source_row(product=""), source_row(product="-")], {"": "5", "-": "5"})
        for row in rows:
            self.assertIsNone(row["reference_price"])
        self.assertIsNone(rows[1]["core_cost"])
        self.assertIsNone(rows[2]["core_cost"])
        self.assertEqual(service.summarize(rows)["paid_amount"], 300)
        self.assertEqual(service.summarize(rows)["identified_customer_rate"], 66.67)

    def test_reference_separates_type_and_online_offline(self):
        rows = self.enriched([
            source_row(amount="100"), source_row(customer="B", amount="200", channel="线上 C"),
            source_row(customer="C", amount="300", product_type="小样"),
        ])
        self.assertEqual([row["reference_price"] for row in rows], [Decimal("10"), Decimal("20"), Decimal("30")])

    def test_reference_never_mixes_conflicting_brands_for_same_sku(self):
        rows = self.enriched([source_row(amount="100"), source_row(customer="B", amount="900", brand="品牌 B")])
        self.assertEqual([row["reference_price"] for row in rows], [Decimal("10"), Decimal("90")])
        previous = self.enriched([source_row(amount="100")], previous=[source_row(amount="900", brand="品牌 B")],
                                 params=service.AnalysisParams(view="prices", benchmark="previous"))
        self.assertIsNone(previous[0]["reference_price"])

    def test_peers_excludes_same_customer_across_channels_and_other_salespeople(self):
        params = service.AnalysisParams(view="prices", benchmark="peers", customer_keyword="客户 A")
        rows = self.enriched([
            source_row(amount="100"), source_row(amount="200", channel="线下 A2"),
            source_row(customer="B", quantity="20", amount="300"),
            source_row(customer="C", quantity="50", amount="5000", channel="线下 B"),
        ], params=params)
        self.assertEqual(len(rows), 2)
        self.assertEqual([row["reference_price"] for row in rows], [Decimal("15"), Decimal("15")])
        self.assertEqual(rows[0]["reference_quantity"], Decimal("20"))

    def test_peers_only_one_customer_returns_no_reference(self):
        rows = self.enriched([source_row()], params=service.AnalysisParams(view="prices", benchmark="peers"))
        self.assertIsNone(rows[0]["reference_price"])

    def test_company_reference_excludes_unidentified_clients(self):
        rows = self.enriched([source_row(amount="100"), source_row(customer="", amount="900")])
        self.assertEqual(rows[0]["reference_price"], Decimal("10"))
        self.assertIsNone(rows[1]["reference_price"])

    def test_unknown_channel_reuses_name_based_online_classifier(self):
        params = service.AnalysisParams(channel_type="online")
        rows = service.prepare_rows([source_row(channel="京东新渠道"), source_row(channel="未知线下")], DIMENSIONS, params)
        self.assertEqual([row["channel"] for row in rows], ["京东新渠道"])

    def test_previous_matches_owner_customer_sku_and_channel_type(self):
        rows = self.enriched([source_row(amount="120")], previous=[
            source_row(amount="80"), source_row(customer="B", amount="300"),
            source_row(amount="400", channel="线下 B"), source_row(amount="600", product="SKU-2"),
        ], params=service.AnalysisParams(view="prices", benchmark="previous"))
        self.assertEqual(rows[0]["reference_price"], Decimal("8"))
        self.assertEqual(rows[0]["price_delta"], Decimal("4"))
        self.assertEqual(rows[0]["price_delta_rate"], Decimal("50"))

    def test_zero_reference_price_does_not_divide_zero(self):
        rows = self.enriched([source_row(amount="100")], previous=[source_row(amount="0")],
                             params=service.AnalysisParams(view="prices", benchmark="previous"))
        self.assertEqual(rows[0]["reference_price"], Decimal("0"))
        self.assertEqual(rows[0]["price_delta"], Decimal("10"))
        self.assertIsNone(rows[0]["price_delta_rate"])

    def test_cost_and_direction_filters_affect_metrics_and_breakdown(self):
        raw = [source_row(amount="100"), source_row(customer="B", amount="200"), source_row(product="SKU-2")]
        params = service.AnalysisParams(view="prices", price_direction="below", cost_status="priced")
        rows = self.enriched(raw, {"SKU-1": "8"}, params)
        self.assertEqual([row["customer_code"] for row in rows], ["A"])
        self.assertEqual(service.summarize(rows)["paid_amount"], 100)
        self.assertEqual(service.breakdown(rows, "customer")[0]["key"], "A")
        self.assertEqual(service.breakdown(rows, "product")[0]["key"], "SKU-1")

    def test_free_goods_are_below_cost_even_zero_sales_amount(self):
        rows = self.enriched([source_row(amount="0")], {"SKU-1": "5"},
                             service.AnalysisParams(view="profit", cost_status="loss"))
        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0]["below_cost"])
        self.assertEqual(service.summarize(rows)["estimated_profit"], -50)

    def test_loss_filter_is_group_average_not_transaction_profit(self):
        rows = self.enriched([source_row(amount="40"), source_row(customer="B", amount="100")], {"SKU-1": "6"},
                             service.AnalysisParams(cost_status="loss"))
        self.assertEqual([row["customer_code"] for row in rows], ["A"])
        self.assertEqual(service.summarize(rows)["loss_sales_amount"], 40)

    def test_sort_null_last_both_directions_stable_ties(self):
        rows = self.enriched([source_row(customer="B"), source_row(customer="A"), source_row(product="SKU-2")], {"SKU-1": "5"})
        for order in ("asc", "desc"):
            sorted_rows = service.sort_rows(rows, service.AnalysisParams(sort_by="estimated_profit", sort_order=order))
            self.assertEqual([row["customer_code"] for row in sorted_rows], ["A", "B", "A"])
            self.assertIsNone(sorted_rows[-1]["estimated_profit"])
        with self.assertRaises(HTTPException):
            service.sort_rows(rows, service.AnalysisParams(sort_by="arbitrary_sql"))

    def test_decimal_rounding_is_half_up(self):
        self.assertEqual(service.number(Decimal("1.005")), 1.01)
        self.assertEqual(service.number(Decimal("-1.005")), -1.01)

    def test_sqlite_sum_tails_never_create_false_price_direction(self):
        raw = [
            source_row(customer="A", quantity="3", amount="0.30000000000000004"),
            source_row(customer="B", quantity="7.000000000000001", amount="0.7000000000000001"),
        ]
        rows = self.enriched(raw)
        self.assertEqual([row["price_delta"] for row in rows], [Decimal("0"), Decimal("0")])
        for direction in ("above", "below"):
            self.assertEqual(self.enriched(raw, params=service.AnalysisParams(view="prices", price_direction=direction)), [])

    def test_genuine_subcent_weighted_price_difference_is_not_erased(self):
        rows = self.enriched([source_row(quantity="3", amount="0.30"), source_row(customer="B", quantity="3", amount="0.31")],
                             params=service.AnalysisParams(view="prices", price_direction="below"))
        self.assertEqual(len(rows), 1)
        self.assertLess(rows[0]["price_delta"], 0)

    def test_period_and_equal_length_previous(self):
        batch = SimpleNamespace(source_start_date=date(2026, 7, 10), source_end_date=date(2026, 10, 8), reconciliation={})
        self.assertEqual(service.resolve_period(service.AnalysisParams(), batch), (
            date(2026, 10, 1), date(2026, 10, 8), date(2026, 9, 23), date(2026, 9, 30),
        ))
        invalid = [
            service.AnalysisParams(range="custom"), service.AnalysisParams(start_date=date(2026, 9, 1)),
            service.AnalysisParams(start_date=date(2026, 10, 8), end_date=date(2026, 10, 1)),
            service.AnalysisParams(start_date=date(2026, 10, 1), end_date=date(2026, 10, 9)),
            service.AnalysisParams(start_date=date(2026, 6, 1), end_date=date(2026, 10, 8)),
            service.AnalysisParams(start_date=date(2026, 7, 10), end_date=date(2026, 8, 1), benchmark="previous"),
        ]
        for params in invalid:
            with self.subTest(params=params), self.assertRaises(HTTPException) as error:
                service.resolve_period(params, batch)
            self.assertEqual(error.exception.status_code, 422)

    def test_sql_binds_filters_and_never_reads_transaction_source(self):
        params = service.AnalysisParams(brand=("O'Reilly",), channel=("渠道 A",), keyword="%_'", product_type=("小样",))
        sql, values = service.aggregation_sql(params, "v1", date(2026, 10, 1), date(2026, 10, 8))
        self.assertIn("FROM `ads_sales_profit_daily`", sql)
        self.assertNotIn("销售单明细账", sql)
        self.assertNotIn("O'Reilly", sql)
        self.assertEqual(values["brand_0"], "O'Reilly")
        self.assertEqual(values["keyword"], "%_'")
        self.assertIn("BETWEEN :start_date AND :end_date", sql)
        self.assertNotIn("owner", values)


class SalesProfitApiTest(unittest.TestCase):
    def setUp(self):
        service._raw_cache.clear()
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        self.db = Session(self.engine)
        AdsPublishBatch.__table__.create(self.engine)
        self.db.execute(text("""CREATE TABLE ads_sales_profit_daily (
            data_version TEXT, item_id INTEGER, sales_date DATE, brand TEXT, channel TEXT,
            customer_code TEXT, customer_name TEXT, product_code TEXT, product_name TEXT, product_type TEXT,
            quantity NUMERIC, paid_amount NUMERIC, sale_quantity NUMERIC, sale_amount NUMERIC,
            return_quantity NUMERIC, return_amount NUMERIC, source_rows INTEGER
        )"""))
        self.db.execute(text("CREATE TABLE `渠道列表` (`渠道名称` TEXT, `负责人` TEXT, `分类` TEXT, `线上平台` TEXT)"))
        self.db.execute(text("INSERT INTO `渠道列表` VALUES ('线下 A', '销售 A', '销售部渠道', NULL), ('线下 B', '销售 B', '销售部渠道', NULL)"))
        self.db.add(AdsPublishBatch(
            data_version="profit-test", dataset="sales_profit", status="ready",
            source_start_date=date(2026, 7, 10), source_end_date=date(2026, 10, 8),
            created_at=datetime(2026, 10, 8, 9), finished_at=datetime(2026, 10, 8, 10),
            published_at=datetime(2026, 10, 8, 10), daily_row_count=12,
            reconciliation={"source_max_updated_at": "2026-10-08T09:44:17"},
        ))
        for index in range(12):
            row = source_row(customer=str(index), product=f"{index:05d}", channel="线下 A" if index < 6 else "线下 B")
            row.update({"data_version": "profit-test", "item_id": index, "sales_date": "2026-10-08", "source_rows": 1})
            self.db.execute(text(f"INSERT INTO ads_sales_profit_daily ({', '.join(row)}) VALUES ({', '.join(':' + key for key in row)})"), row)
        self.db.commit()
        self.permissions = {"sales.profit.view", "data.export"}
        self.user = SimpleNamespace(role=SimpleNamespace(is_active=True, permissions=[SimpleNamespace(code=code) for code in self.permissions]))
        app = FastAPI()
        app.include_router(router_module.router, prefix="/api/v1", dependencies=[Depends(require_permission("sales.profit.view"))])
        app.dependency_overrides[get_sales_profit_db] = lambda: self.db
        app.dependency_overrides[get_ods_db] = lambda: self.db
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_user] = lambda: self.user
        self.client = TestClient(app)
        self.price_patch = patch.object(router_module, "current_prices", return_value={f"{index:05d}": Decimal("6") for index in range(8)})
        self.price_patch.start()

    def tearDown(self):
        self.price_patch.stop()
        self.client.close()
        self.db.close()
        self.engine.dispose()
        service._raw_cache.clear()

    def get(self, **params):
        return self.client.get("/api/v1/sales/profit-analysis", params=params)

    def test_default_pagination_and_full_summary(self):
        response = self.get(page_size=10)
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()["data"]
        self.assertEqual(len(data["rows"]), 10)
        self.assertEqual(data["pagination"]["total"], 12)
        self.assertEqual(data["summary"]["paid_amount"], 1200)
        self.assertEqual(data["summary"]["estimated_profit"], 320)
        self.assertEqual(data["summary"]["cost_coverage_rate"], 66.67)
        self.assertEqual(data["updated_at"], "2026-10-08T09:44:17")
        self.assertEqual(data["options"]["owners"], ["销售 A", "销售 B"])
        self.assertEqual(data["start_date"], "2026-10-01")
        self.assertEqual(self.get(page=2, page_size=10).json()["data"]["pagination"]["total"], 12)

    def test_combined_filters_and_cache_scope(self):
        first = self.get(owner="销售 A", customer_keyword="客户 1", brand="品牌 A", channel="线下 A", product_type="正装")
        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(first.json()["data"]["summary"]["paid_amount"], 100)
        second = self.get(owner="销售 B", cost_status="missing", keyword="000", sort_by="product_code", sort_order="asc")
        self.assertEqual(second.status_code, 200, second.text)
        self.assertEqual(second.json()["data"]["pagination"]["total"], 4)
        self.assertEqual(second.json()["data"]["rows"][0]["product_code"], "00008")

    def test_default_types_and_all_sentinel(self):
        self.db.execute(text("UPDATE ads_sales_profit_daily SET product_type='其他' WHERE item_id=11"))
        self.db.commit()
        self.assertEqual(self.get().json()["data"]["pagination"]["total"], 11)
        self.assertEqual(self.get(product_type="__all__").json()["data"]["pagination"]["total"], 12)

    def test_export_all_rows_same_filters_and_numeric_code_string(self):
        response = self.client.get("/api/v1/sales/profit-analysis/export", params={"owner": "销售 A", "page_size": 10})
        self.assertEqual(response.status_code, 200, response.text)
        sheet = load_workbook(BytesIO(response.content)).worksheets[0]
        self.assertEqual(sheet.max_row, 7)
        self.assertEqual(sheet.cell(2, 4).value, "00000")
        self.assertEqual(sheet.cell(2, 4).data_type, "s")
        self.assertEqual(sheet.cell(2, 19).value, 40)
        self.assertEqual(sheet.cell(2, 19).number_format, "#,##0.00")
        self.assertEqual(sheet.freeze_panes, "F2")
        self.assertEqual(sheet.auto_filter.ref, "A1:Y7")

    def test_return_exclusion_matches_detail_summary_and_export(self):
        self.db.execute(text("UPDATE ads_sales_profit_daily SET quantity=-90,return_quantity=100,return_amount=0 WHERE item_id=0"))
        self.db.commit()
        data = self.get(keyword='00000').json()['data']
        self.assertEqual(data['rows'][0]['quantity'], -90)
        self.assertEqual(data['rows'][0]['estimated_profit'], 40)
        self.assertEqual(data['summary']['estimated_profit'], 40)
        self.assertEqual(data['breakdown'][0]['sale_quantity'], 10)
        self.assertEqual(data['breakdown'][0]['sale_amount'], 100)
        response = self.client.get('/api/v1/sales/profit-analysis/export', params={'keyword': '00000'})
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.worksheets[0].cell(2, 18).value, 60)
        self.assertEqual(book.worksheets[0].cell(2, 19).value, 40)
        self.assertIn('退货红冲不参与毛利', book['统计口径'].cell(5, 2).value)
        book.close()

    def test_export_preserves_literal_formula_text(self):
        self.db.execute(text("UPDATE ads_sales_profit_daily SET product_name='=HYPERLINK(\"https://invalid\")' WHERE item_id=0"))
        self.db.commit()
        response = self.client.get("/api/v1/sales/profit-analysis/export", params={"keyword": "00000"})
        sheet = load_workbook(BytesIO(response.content)).worksheets[0]
        self.assertEqual(sheet.cell(2, 5).data_type, "s")
        self.assertEqual(sheet.cell(2, 5).value, '=HYPERLINK("https://invalid")')

    def test_route_validation(self):
        for params in ({"sort_by": "injection"}, {"view": "other"}, {"page_size": 1000}, {"channel_type": "all_unchecked"}, {"group_by": "user"}):
            self.assertEqual(self.get(**params).status_code, 422)
        self.assertEqual(self.get(range="custom").status_code, 422)
        self.assertEqual(self.get(start_date="2026-10-01", end_date="2026-10-09").status_code, 422)
        self.assertEqual(self.get(brand="x" * 256).status_code, 422)

    def test_profit_ignores_hidden_price_controls(self):
        response = self.get(view="profit", benchmark="previous", price_direction="below")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["data"]["pagination"]["total"], 12)

    def test_previous_price_requires_coverage(self):
        response = self.get(view="prices", benchmark="previous", start_date="2026-07-10", end_date="2026-08-01")
        self.assertEqual(response.status_code, 422)

    def test_new_module_permission_and_export_permission(self):
        self.user.role.permissions = [SimpleNamespace(code="sales.view")]
        self.assertEqual(self.get().status_code, 403)
        self.assertEqual(self.client.get("/api/v1/sales/profit-analysis/export").status_code, 403)
        self.user.role.permissions = [SimpleNamespace(code="sales.profit.view")]
        self.assertEqual(self.get().status_code, 200)
        self.assertEqual(self.client.get("/api/v1/sales/profit-analysis/export").status_code, 403)

    def test_no_batch_returns_503_without_ods_fact_fallback(self):
        self.db.query(AdsPublishBatch).delete()
        self.db.commit()
        self.assertEqual(self.get().status_code, 503)

    def test_unauthed_denied(self):
        app = FastAPI()
        app.include_router(router_module.router, dependencies=[Depends(require_permission("sales.profit.view"))])
        client = TestClient(app)
        self.assertEqual(client.get("/sales/profit-analysis").status_code, 401)
        client.close()

    def test_missing_table_returns_503_without_fallback(self):
        self.db.execute(text("DROP TABLE ads_sales_profit_daily"))
        self.db.commit()
        self.assertEqual(self.get().status_code, 503)

    def test_result_guard_does_not_silently_truncate_export(self):
        with patch.object(service, "MAX_GROUPED_ROWS", 5):
            response = self.client.get("/api/v1/sales/profit-analysis/export")
        self.assertEqual(response.status_code, 422)

    def test_owner_filter_runs_before_group_limit_and_empty_period_is_valid(self):
        self.db.execute(text("UPDATE ads_sales_profit_daily SET sales_date='2026-09-15'"))
        self.db.commit()
        query = {"start_date": "2026-09-01", "end_date": "2026-09-30"}
        with patch.object(service, "MAX_GROUPED_ROWS", 6):
            selected = self.get(owner="销售 A", **query)
            self.assertEqual(selected.status_code, 200, selected.text)
            self.assertEqual(selected.json()["data"]["pagination"]["total"], 6)
            self.assertEqual(self.get(**query).status_code, 422)
            missing = self.get(owner="没有销售记录", **query)
            self.assertEqual(missing.status_code, 200, missing.text)
            self.assertEqual(missing.json()["data"]["pagination"]["total"], 0)
            self.assertEqual(missing.json()["data"]["start_date"], "2026-09-01")


if __name__ == "__main__":
    unittest.main()
