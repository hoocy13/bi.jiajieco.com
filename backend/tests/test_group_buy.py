import unittest
from unittest.mock import Mock
from datetime import datetime
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import xlsxwriter
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api.deps import get_current_user
from app.api.routers import group_buy
from app.db.session import get_db
from app.db.ods import get_ods_db
from app.models.group_buy import GroupBuyBatch, GroupBuyChange
from app.services import group_buy as service


def workbook(kind, codes=("A", "B"), negative=False):
    out = BytesIO()
    book = xlsxwriter.Workbook(out, {"in_memory": True})
    sheet = book.add_worksheet()
    sheet.write_row(7, 0, ["国际条形码", "产品代码", "产品名称", "价格", "ABC", "PCB"])
    if kind == "order":
        sheet.write(2, 6, "柜台甲")
    else:
        sheet.write_row(1, 7, ["柜台甲", "柜台甲", "吉客云"])
        sheet.write_row(2, 7, ["客户甲", "客户甲", "客户乙"])
    for r, code in enumerate(codes, 8):
        sheet.write_row(r, 0, ["001234", code, "商品"+code, 12.35, "OS" if r == 8 else "A", 20])
        if kind == "order":
            sheet.write(r, 6, -1 if negative else 10)
        else:
            sheet.write_row(r, 7, [5, 5, 2])
    book.close()
    return out.getvalue()


class GroupBuyTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        GroupBuyBatch.__table__.create(self.engine)
        GroupBuyChange.__table__.create(self.engine)
        self.db = Session(self.engine)
        self.user = SimpleNamespace(username="tester", role=SimpleNamespace(is_active=True, permissions=[SimpleNamespace(code="group_buy.manage")]))
        app = FastAPI()
        self.app = app
        app.include_router(group_buy.router, prefix="/api/v1")
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_user] = lambda: self.user
        self.client = TestClient(app)
        self.base = "/api/v1/group-buy/batches"
        r = self.client.post(self.base, json={"month": "2026-10", "brand": "兰蔻", "name": "第一批"})
        self.assertEqual(r.status_code, 200, r.text)
        self.id = r.json()["data"]["id"]
        self.url = self.base + f"/{self.id}"

    def tearDown(self):
        self.client.close()
        self.db.close()
        self.engine.dispose()

    def upload(self, stage="apply", revision=1, order=None, allocation=None):
        return self.client.post(self.url+"/import/"+stage, data={"revision": revision}, files={
            "order_file": ("下单.xlsx", order or workbook("order")),
            "allocation_file": ("分配.xlsx", allocation or workbook("allocation")),
        })

    def table(self, **params):
        r = self.client.get(self.url, params=params)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()["data"]

    def test_import_preview_is_read_only_and_preserves_duplicates_unknown_quota(self):
        r = self.upload("preview")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["data"]["duplicate_columns"], 1)
        self.assertEqual(self.table()["pagination"]["total"], 0)
        r = self.upload()
        self.assertEqual(r.status_code, 200, r.text)
        d = self.table()
        self.assertEqual(d["summary"]["ordered"], 20)
        self.assertEqual(d["summary"]["allocated"], 24)
        self.assertEqual(d["summary"]["missing_count"], 2)
        self.assertEqual(len({c["id"] for c in d["columns"]}), 3)
        self.assertEqual(d["rows"][0]["barcode"], "001234")
        self.assertEqual(self.client.post(self.url+"/confirm", json={"revision":2}).status_code, 422)
        self.assertEqual(self.upload(revision=2).status_code, 422)

    def test_filters_sort_pagination_and_negative_rejection(self):
        self.assertEqual(self.upload(order=workbook("order", negative=True)).status_code, 422)
        codes = tuple(f"P{i:02}" for i in range(15))
        self.assertEqual(self.upload(order=workbook("order", codes), allocation=workbook("allocation", codes)).status_code, 200)
        d = self.table()
        source = d["sources"][0]["id"]
        filtered = self.table(only_os=True, keyword="P00", source=source, customer="客户甲")
        self.assertEqual(filtered["pagination"]["total"], 1)
        self.assertEqual(len(filtered["columns"]), 2)
        self.assertEqual(filtered["summary"]["allocated"], 12)
        page = self.table(page=2, page_size=10, sort_by="code", sort_order="desc")
        self.assertEqual([p["code"] for p in page["rows"]], ["P04", "P03", "P02", "P01", "P00"])
        full = self.table(all_rows=True, page=2, page_size=10, sort_by="code", sort_order="desc")
        self.assertEqual(len(full["rows"]), 15)
        self.assertEqual(full["rows"][-1]["code"], "P00")
        self.assertEqual(full["pagination"], {"page": 1, "page_size": 15, "total": 15})
        combined = self.table(all_rows=True, keyword="P0", source=source, customer="客户甲", only_os=True)
        self.assertEqual([r["code"] for r in combined["rows"]], ["P00"])
        self.assertEqual(combined["summary"]["allocated"], 12)

    def test_atomic_save_conflict_and_audit(self):
        self.upload()
        d = self.table()
        p, c = d["rows"][0], d["columns"][0]
        cell = {"product_id":p["id"], "kind":"allocation", "column_id":c["id"], "quantity":3}
        r = self.client.patch(self.url+"/cells", json={"revision":2,"cells":[cell]})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(self.client.patch(self.url+"/cells", json={"revision":2,"cells":[cell]}).status_code, 409)
        self.assertEqual(self.table()["rows"][0]["allocations"][c["id"]], 3)
        invalid = {**cell,"column_id":"missing"}
        r = self.client.patch(self.url+"/cells", json={"revision":3,"cells":[{**cell,"quantity":1},invalid]})
        self.assertEqual(r.status_code, 422, r.text)
        self.assertEqual(self.table()["rows"][0]["allocations"][c["id"]], 3)
        history = self.client.get(self.url+"/history").json()["data"]["rows"]
        self.assertEqual(history[0]["operator"], "tester")
        self.assertEqual(history[0]["changes"]["cells"][0]["before"], 5)
        for qty in [-1, 1.5, "3"]:
            self.assertEqual(self.client.patch(self.url+"/cells",json={"revision":3,"cells":[{**cell,"quantity":qty}]}).status_code,422)

    def test_quota_confirm_lock_reopen_copy(self):
        self.upload()
        d = self.table()
        revision = 2
        for p in d["rows"]:
            r = self.client.post(self.url+"/quotas",json={"revision":revision,"product_id":p["id"],"source":"吉客云","quantity":2,"mode":"set","note":"核对现货"})
            self.assertEqual(r.status_code,200,r.text)
            revision+=1
        self.assertEqual(self.client.post(self.url+"/confirm",json={"revision":revision}).status_code,200)
        revision+=1
        r = self.client.post(self.url+"/columns",json={"revision":revision,"source":"柜台甲","customer":"新客户"})
        self.assertEqual(r.status_code,409)
        self.assertEqual(self.client.post(self.url+"/reopen",json={"revision":revision}).status_code,200)
        r=self.client.post(self.base,json={"month":"2026-11","brand":"兰蔻","name":"第一批","copy_from":self.id})
        self.assertEqual(r.status_code,200,r.text)
        copied=self.client.get(self.base+f'/{r.json()["data"]["id"]}').json()["data"]
        self.assertEqual(copied["summary"]["allocated"],0)
        self.assertEqual(copied["summary"]["known_quota"],0)
        self.assertEqual(copied["rows"][0]["price"],"12.35")

    def test_export_roundtrip_and_quota_ledger(self):
        self.upload()
        d=self.table();p=d["rows"][0]
        self.client.post(self.url+"/quotas",json={"revision":2,"product_id":p["id"],"source":"吉客云","quantity":5,"mode":"set","note":"现货核对"})
        exported={}
        for kind in ("order","allocation"):
            r=self.client.get(self.url+"/export",params={"kind":kind})
            self.assertEqual(r.status_code,200)
            exported[kind]=service.parse_workbook(r.content,kind)
            book=load_workbook(BytesIO(r.content),data_only=True)
            self.assertIn("分配额度",book.sheetnames)
            self.assertEqual(book.active["A9"].value,"001234")
            self.assertEqual(book.active["D9"].value,12.35)
            book.close()
        result=service.merge_import(service.empty_document(),**exported)
        self.assertEqual(service.summary(result)["missing_count"],1)
        s=next(s for s in result["sources"] if s["name"]=="吉客云")
        self.assertEqual(result["products"][0]["quotas"][s["id"]],5)
        self.assertIsNone(result["products"][1]["quotas"][s["id"]])

    def test_permission_on_every_operation(self):
        self.user.role.permissions=[]
        for method,path,kwargs in [("get",self.base,{}),("get",self.url,{}),("get",self.url+"/export",{}),
                                   ("delete",self.url,{"json":{"revision":1}}),
                                   ("patch",self.url,{"json":{"revision":1,"name":"无权限"}}),
                                   ("post",self.url+"/restore",{"json":{"revision":1}}),
                                   ("get",self.base,{"params":{"deleted":True}}),
                                   ("post",self.url+"/confirm",{"json":{"revision":1}}),
                                   ("post",self.url+"/import/preview",{"data":{"revision":1}})]:
            r=self.client.request(method.upper(),path,**kwargs)
            self.assertEqual(r.status_code,403,r.text)

    def test_two_workbooks_must_match_and_initial_prices_are_snapshots(self):
        self.assertEqual(self.upload(allocation=workbook("allocation",("A","C"))).status_code,422)
        self.assertEqual(self.table()["pagination"]["total"],0)
        self.assertEqual(self.upload().status_code,200)
        self.assertEqual(self.client.post(self.base,json={"month":"2026-10","brand":"兰蔻","name":"第一批"}).status_code,409)

    def test_product_edit_and_quota_adjustment_keep_other_batches_unchanged(self):
        self.upload()
        d = self.table(); p=d["rows"][0]
        copy=self.client.post(self.base,json={"month":"2026-11","brand":"兰蔻","name":"副本","copy_from":self.id}).json()["data"]
        body={k:p[k] for k in ("code","name","barcode","price","abc","pcb")}
        r=self.client.patch(self.url+f'/products/{p["id"]}',json={**body,"price":"15.50","revision":2})
        self.assertEqual(r.status_code,200,r.text)
        copied=self.client.get(self.base+f'/{copy["id"]}').json()["data"]
        self.assertEqual(copied["rows"][0]["price"],"12.35")
        self.assertEqual(self.table()["rows"][0]["price"],"15.50")
        src=next(s for s in d["sources"] if s["counter"])
        r=self.client.post(self.url+"/quotas",json={"revision":3,"product_id":p["id"],"source":src["name"],"quantity":3,"mode":"add","note":"追加"})
        self.assertEqual(r.status_code,200)
        r=self.client.patch(self.url+"/cells",json={"revision":4,"cells":[{"product_id":p["id"],"kind":"order","column_id":src["id"],"quantity":12}]})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(self.table()["rows"][0]["quotas"][src["id"]],15)
        r=self.client.post(self.url+"/quotas",json={"revision":5,"product_id":p["id"],"source":"吉客云","quantity":3,"mode":"add","note":"未知额度不能直接追加"})
        self.assertEqual(r.status_code,422)

    def test_batch_recycle_preserves_data_and_blocks_normal_access(self):
        self.upload()
        before = self.table()
        self.assertEqual(self.client.request("DELETE", self.url, json={"revision": 1}).status_code, 409)
        r = self.client.request("DELETE", self.url, json={"revision": 2})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["data"]["status"], "deleted")
        self.assertEqual(self.client.get(self.base).json()["data"]["rows"], [])
        recycled = self.client.get(self.base, params={"month":"2026-10", "brand":"兰蔻", "deleted":True}).json()["data"]["rows"]
        self.assertEqual([r["id"] for r in recycled], [self.id])
        self.assertEqual(recycled[0]["summary"]["allocated"], before["summary"]["allocated"])
        for path in (self.url, self.url+"/export"):
            self.assertEqual(self.client.get(path).status_code, 404)
        self.assertEqual(self.client.post(self.url+"/reopen", json={"revision":3}).status_code,404)
        p, c = before["rows"][0], before["columns"][0]
        self.assertEqual(self.client.patch(self.url+"/cells",json={"revision":3,"cells":[{"product_id":p["id"],"column_id":c["id"],"kind":"allocation","quantity":1}]}).status_code,404)
        history = self.client.get(self.url+"/history").json()["data"]["rows"]
        self.assertEqual(history[0]["action"], "delete")
        self.assertEqual(history[0]["operator"], "tester")
        self.assertEqual(self.client.post(self.url+"/restore",json={"revision":2}).status_code,409)
        restored = self.client.post(self.url+"/restore",json={"revision":3})
        self.assertEqual(restored.status_code,200,restored.text)
        self.assertEqual(restored.json()["data"]["status"],"draft")
        after = self.table()
        self.assertEqual(after["rows"], before["rows"])
        self.assertEqual(after["sources"],before["sources"])
        self.assertEqual(after["all_columns"],before["all_columns"])
        self.assertEqual(self.client.post(self.url+"/restore",json={"revision":4}).status_code,409)

    def test_restore_confirmed_batch_keeps_lock_and_rename_is_atomic(self):
        self.db.get(GroupBuyBatch, self.id).status = "confirmed"
        self.db.commit()
        r = self.client.request("DELETE",self.url,json={"revision":1})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(self.client.post(self.url+"/restore",json={"revision":2}).json()["data"]["status"],"confirmed")
        renamed=self.client.patch(self.url,json={"revision":3,"name":"  已改名  "})
        self.assertEqual(renamed.status_code,200,renamed.text)
        self.assertEqual(renamed.json()["data"]["name"],"已改名")
        self.assertEqual(renamed.json()["data"]["status"],"confirmed")
        self.assertEqual(self.client.patch(self.url,json={"revision":3,"name":"旧版本"}).status_code,409)
        self.assertEqual(self.client.patch(self.url,json={"revision":4,"name":"   "}).status_code,422)
        self.client.post(self.base,json={"month":"2026-10","brand":"兰蔻","name":"重名"})
        conflict=self.client.patch(self.url,json={"revision":4,"name":"重名"})
        self.assertEqual(conflict.status_code,409)
        self.assertEqual(self.table()["batch"]["name"],"已改名")
        self.assertEqual(self.table()["batch"]["revision"],4)

    def test_remove_customer_column_retains_orders_quotas_and_duplicate_sibling(self):
        self.upload()
        before = self.table()
        c = before["columns"][0]
        endpoint = self.url + f'/columns/{c["id"]}'
        self.assertEqual(self.client.request("DELETE",endpoint,json={"revision":1}).status_code,409)
        self.assertEqual(self.client.request("DELETE",endpoint,json={"revision":2}).status_code,200)
        after = self.table()
        self.assertEqual(len(after["columns"]),2)
        self.assertEqual(after["summary"]["allocated"],14)
        for old, new in zip(before["rows"],after["rows"]):
            self.assertEqual(new["orders"],old["orders"])
            self.assertEqual(new["quotas"],old["quotas"])
            self.assertNotIn(c["id"],new["allocations"])
        audit=self.client.get(self.url+"/history").json()["data"]["rows"][0]
        self.assertEqual(audit["changes"]["removed"],c)
        self.assertEqual(sum(audit["changes"]["allocations"].values()),10)
        self.assertEqual(self.client.request("DELETE",endpoint,json={"revision":3}).status_code,404)

    def test_remove_source_cascades_only_its_orders_quotas_allocations_and_exports(self):
        self.upload()
        before=self.table(); source=before["sources"][0]
        copied=self.client.post(self.base,json={"month":"2026-11","brand":"兰蔻","name":"副本","copy_from":self.id}).json()["data"]
        self.assertEqual(self.client.request("DELETE",self.url+f'/sources/{source["id"]}',json={"revision":2}).status_code,200)
        after=self.table()
        self.assertEqual(after["summary"]["ordered"],0)
        self.assertEqual(after["summary"]["allocated"],4)
        self.assertEqual(len(after["columns"]),1)
        for p in after["rows"]:
            self.assertNotIn(source["id"],p["orders"])
            self.assertNotIn(source["id"],p["quotas"])
            self.assertEqual(set(p["allocations"]),{after["columns"][0]["id"]})
        audit=self.client.get(self.url+"/history").json()["data"]["rows"][0]["changes"]
        self.assertEqual(len(audit["columns"]),2)
        self.assertEqual(sum(p["orders"].get(source["id"],0) for p in audit["products"]),20)
        other=self.client.get(self.base+f'/{copied["id"]}').json()["data"]
        self.assertEqual(len(other["sources"]),2)
        for kind in ("order","allocation"):
            self.assertEqual(self.client.get(self.url+"/export",params={"kind":kind}).status_code,200)

    def test_product_delete_removes_both_tables_and_confirmed_batch_blocks_structure_changes(self):
        self.upload()
        before=self.table();p=before["rows"][0]
        self.assertEqual(self.client.request("DELETE",self.url+f'/products/{p["id"]}',json={"revision":2}).status_code,200)
        self.assertEqual(self.table()["summary"]["ordered"],10)
        self.assertEqual(self.table()["summary"]["allocated"],12)
        audit=self.client.get(self.url+"/history").json()["data"]["rows"][0]
        self.assertEqual(audit["changes"]["removed"]["allocations"],p["allocations"])
        self.db.get(GroupBuyBatch,self.id).status="confirmed";self.db.commit()
        for collection,item in (("products",before["rows"][1]),("columns",before["columns"][0]),("sources",before["sources"][0])):
            self.assertEqual(self.client.request("DELETE",self.url+f'/{collection}/{item["id"]}',json={"revision":3}).status_code,409)

    def test_rename_source_customer_preserves_ids_values_and_rejects_collisions(self):
        self.upload()
        d=self.table();source=d["sources"][0];c=d["columns"][0]
        self.assertEqual(self.client.patch(self.url+f'/sources/{source["id"]}',json={"revision":2,"name":"柜台新名"}).status_code,200)
        self.assertEqual(self.client.patch(self.url+f'/columns/{c["id"]}',json={"revision":3,"name":"客户新名"}).status_code,200)
        after=self.table()
        self.assertEqual(after["rows"],d["rows"])
        self.assertEqual(after["sources"][0]["id"],source["id"])
        self.assertEqual(after["columns"][0]["customer"],"客户新名")
        self.assertEqual(self.client.patch(self.url+f'/sources/{source["id"]}',json={"revision":4,"name":"吉客云"}).status_code,422)
        self.assertEqual(self.client.patch(self.url+f'/columns/{c["id"]}',json={"revision":4,"name":"客户甲"}).status_code,422)
        self.assertEqual(self.table()["batch"]["revision"],4)
        r=self.client.post(self.url+"/columns",json={"revision":4,"source":"独立柜台","counter":True})
        self.assertEqual(r.status_code,200,r.text)
        final=self.table()
        self.assertEqual(len(final["all_columns"]),3)
        self.assertTrue(final["sources"][-1]["counter"])
        self.user.role.permissions=[]
        for method,path,body in (("DELETE",f'/products/{d["rows"][0]["id"]}',{"revision":5}),
                                 ("DELETE",f'/columns/{c["id"]}',{"revision":5}),
                                 ("DELETE",f'/sources/{source["id"]}',{"revision":5}),
                                 ("PATCH",f'/columns/{c["id"]}',{"revision":5,"name":"禁止"}),
                                 ("PATCH",f'/sources/{source["id"]}',{"revision":5,"name":"禁止"})):
            self.assertEqual(self.client.request(method,self.url+path,json=body).status_code,403)

    def test_catalog_exact_lookup_priority_fallback_conflicts_and_permission(self):
        product={"code":"SKU-001","name":"标准商品名","barcode":"0012345678901","updated_at":datetime(2026,10,22,9)}
        ods=Mock()
        def result(rows):
            value=Mock();value.mappings.return_value.all.return_value=rows;return value
        self.app.dependency_overrides[get_ods_db]=lambda:ods
        ods.execute.return_value=result([product])
        endpoint="/api/v1/group-buy/catalog"
        r=self.client.get(endpoint,params={"code":" SKU-001 "})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()["data"]["product"]["barcode"],"0012345678901")
        self.assertEqual(r.json()["data"]["product"]["source"],"总库存查询")
        self.assertEqual(ods.execute.call_count,1)
        sql,params=ods.execute.call_args.args
        self.assertIn("WHERE `货品编号` = :code",str(sql))
        self.assertEqual(params,{"code":"SKU-001"})
        ods.execute.reset_mock();ods.execute.side_effect=[result([]),result([product])]
        fallback=self.client.get(endpoint,params={"code":"SKU-001"}).json()["data"]
        self.assertEqual(fallback["product"]["source"],"分仓库查询")
        ods.execute.side_effect=[result([]),result([])]
        self.assertEqual(self.client.get(endpoint,params={"code":"SKU-001-OTHER"}).json()["data"]["status"],"not_found")
        ods.execute.side_effect=[result([product,{**product,"barcode":"different"}])]
        self.assertEqual(self.client.get(endpoint,params={"code":"SKU-001"}).json()["data"]["status"],"ambiguous")
        ods.execute.side_effect=[result([{**product,"barcode":""}])]
        self.assertEqual(self.client.get(endpoint,params={"code":"SKU-001"}).json()["data"]["status"],"incomplete")
        self.assertEqual(self.client.get(endpoint,params={"code":" "}).status_code,422)
        self.assertEqual(self.table()["batch"]["revision"],1)
        self.user.role.permissions=[]
        ods.execute.reset_mock()
        self.assertEqual(self.client.get(endpoint,params={"code":"SKU-001"}).status_code,403)
        ods.execute.assert_not_called()

    def test_role_startup_preserves_explicit_group_buy_grants_and_revocations(self):
        from app.db.init_db import seed_roles
        from app.models.user import Permission, Role, role_permissions
        for table in (Permission.__table__, Role.__table__, role_permissions):
            table.create(self.engine, checkfirst=True)
        roles=seed_roles(self.db)
        self.db.commit()
        permission=self.db.query(Permission).filter_by(code="group_buy.manage").one()
        for code in ("data_analyst","management","sales"):
            self.assertNotIn(permission,roles[code].permissions)
        self.assertIn(permission,roles["admin"].permissions)
        for code in ("data_analyst","management"):
            roles[code].permissions.append(permission)
        self.db.commit()
        roles=seed_roles(self.db)
        self.db.commit()
        for code in ("data_analyst","management"):
            self.assertIn(permission,roles[code].permissions)
            roles[code].permissions.remove(permission)
        self.db.commit()
        roles=seed_roles(self.db)
        self.db.commit()
        for code in ("data_analyst","management"):
            self.assertNotIn(permission,roles[code].permissions)


if __name__ == "__main__":
    unittest.main()
