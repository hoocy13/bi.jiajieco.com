import unittest
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import Depends, FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
import xlsxwriter

from app.api.deps import get_current_user, require_permission
from app.api.routers import inventory_valuation as router
from app.db.ods import get_ods_db
from app.db.session import get_db
from app.models.core_cost_price import CoreCostPriceRevision
from app.services.core_cost_catalog import calculate_cost


def workbook_bytes(headers, rows):
    output = BytesIO()
    book = xlsxwriter.Workbook(output, {"in_memory": True, "strings_to_formulas": False})
    sheet = book.add_worksheet()
    sheet.write_row(0, 0, headers)
    for i, row in enumerate(rows, 1):
        sheet.write_row(i, 0, row)
    book.close()
    return output.getvalue()


class CoreCostPricesTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        CoreCostPriceRevision.__table__.create(self.engine)
        self.db = Session(self.engine)
        for table in ("总库存查询", "分仓库查询"):
            self.db.execute(text(f"CREATE TABLE `{table}` (`货品编号` TEXT, `货品名称` TEXT, `品牌` TEXT, `货品分类` TEXT, `零售价` NUMERIC, `含税价` NUMERIC)"))
        self.db.execute(text("INSERT INTO `总库存查询` VALUES ('A','商品A','品牌甲','正装',1000,200),('B','商品B','品牌甲','小样',NULL,30),('D','商品D','品牌甲','正装',100,20),('D','商品D','品牌甲','正装',120,25)"))
        self.db.execute(text("INSERT INTO `分仓库查询` VALUES ('A','商品A','品牌甲','正装',999,111),('C','商品C','品牌甲','小样',50,18)"))
        for table in ('总库存查询', '分仓库查询'):
            self.db.execute(text(f'ALTER TABLE `{table}` ADD COLUMN `条码` TEXT'))
            self.db.execute(text(f"UPDATE `{table}` SET `条码` = '123456' WHERE `货品编号` = 'A'"))
        for sku, price in [('A',350),('Z',20)]:
            self.db.add(CoreCostPriceRevision(product_code=sku,product_name=sku,price=price,effective_date=date.today(),source='import',operator='legacy'))
        self.db.commit()
        self.user = SimpleNamespace(username='tester',role=SimpleNamespace(is_active=True,permissions=[SimpleNamespace(code='inventory.view'),SimpleNamespace(code='inventory.price.manage')]))
        app=FastAPI()
        app.include_router(router.router,prefix='/api/v1',dependencies=[Depends(require_permission('inventory.view'))])
        app.dependency_overrides[get_current_user]=lambda:self.user
        app.dependency_overrides[get_db]=lambda:self.db
        app.dependency_overrides[get_ods_db]=lambda:self.db
        self.client=TestClient(app)
        self.base='/api/v1/inventory/valuation/prices'

    def tearDown(self):
        self.client.close();self.db.close();self.engine.dispose()

    def rows(self,params=None):
        r=self.client.get(self.base,params=params)
        self.assertEqual(r.status_code,200,r.text)
        return r.json()['data']

    def preview(self,headers,rows):
        r=self.client.post(self.base+'/import-preview',files={'file':('prices.xlsx',workbook_bytes(headers,rows))})
        self.assertEqual(r.status_code,200,r.text)
        return r.json()['data']

    def test_catalog_includes_unpriced_fallback_orphans_and_conflicts(self):
        rows={r['product_code']:r for r in self.rows()['rows']}
        self.assertEqual(set(rows),{'A','B','C','D','Z'})
        self.assertEqual((rows['A']['retail_price'],rows['A']['tax_price'],rows['A']['price']),(1000,200,350))
        self.assertEqual(rows['A']['pricing_mode'],'fixed')
        self.assertEqual(rows['C']['catalog_source'],'分仓库查询')
        self.assertTrue(rows['D']['retail_price_conflict'])
        self.assertIsNone(rows['D']['retail_price'])
        self.assertIsNone(rows['B']['price'])

    def test_maintenance_status_and_reference_discount(self):
        rows = self.rows({'source': 'priced'})['rows']
        self.assertEqual({r['product_code'] for r in rows}, {'A', 'Z'})
        self.assertEqual(rows[0]['reference_discount'], 3.5)
        self.assertIsNone(rows[0]['discount'])
        self.assertIsNone(rows[1]['reference_discount'])
        missing = self.rows({'source': 'missing', 'brand': '品牌甲'})['rows']
        self.assertEqual({r['product_code'] for r in missing}, {'B', 'C', 'D'})
        self.assertTrue(all(r['price'] is None for r in missing))

    def test_barcode_matching_and_discount_range(self):
        row = self.rows({'keyword': '123456', 'source': 'priced', 'min_discount': 35, 'max_discount': 35})['rows'][0]
        self.assertEqual(row['barcode'], '123456')
        self.assertEqual(row['product_type'], '正装')
        self.assertEqual(self.rows({'min_discount': 36})['rows'], [])
        self.assertEqual(self.client.get(self.base, params={'min_discount': 50, 'max_discount': 40}).status_code, 422)
        self.db.add(CoreCostPriceRevision(product_code='OLD', barcode='123456', product_name='old', price=350, effective_date=date.today(), source='manual', operator='test'))
        self.db.commit()
        row = self.rows({'keyword': 'OLD'})['rows'][0]
        self.assertEqual((row['barcode'], row['product_type'], row['retail_price']), ('123456', '正装', 1000))
        response = self.client.get(self.base + '/export', params={'min_discount': 35, 'max_discount': 35})
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.active.max_row, 3)
        book.close()

    def test_combined_filter_sort_and_pagination(self):
        data=self.rows({'brand':'品牌甲','product_type':'小样','source':'missing','sort_by':'retail_price','sort_order':'desc'})
        self.assertEqual([r['product_code'] for r in data['rows']],['C','B'])
        for i in range(12):
            self.db.execute(text("INSERT INTO `总库存查询` (`货品编号`,`货品名称`,`品牌`,`货品分类`,`零售价`,`含税价`) VALUES (:code,'商品','品牌乙','正装',10,3)"),{'code':f'E{i:02}'})
        self.db.commit()
        data=self.rows({'page':2,'page_size':10})
        self.assertEqual(data['pagination']['total'],17)
        self.assertEqual(len(data['rows']),7)

    def test_discount_uses_retail_and_snapshot_does_not_follow_retail_updates(self):
        r=self.client.post(self.base,json={'product_code':'A','pricing_mode':'discount','discount':3.5})
        self.assertEqual(r.status_code,200,r.text)
        saved=router.current_prices(self.db)['A']
        self.assertEqual((saved.price,saved.base_retail_price,saved.discount),(Decimal('350'),Decimal('1000'),Decimal('3.5')))
        self.db.execute(text("UPDATE `总库存查询` SET `零售价`=1200 WHERE `货品编号`='A'"));self.db.commit()
        row=self.rows({'source':'changed'})['rows'][0]
        self.assertEqual(row['price'],350)
        self.assertTrue(row['base_changed'])
        self.assertEqual(calculate_cost('2.01','5'),Decimal('1.01'))

    def test_discount_validation_and_stale_basis(self):
        for body,code in [({'product_code':'B'},422),({'product_code':'D'},422),({'product_code':'A','base_retail_price':900},409),({'product_code':'A','price':349},422)]:
            r=self.client.post(self.base,json={'pricing_mode':'discount','discount':3.5,**body})
            self.assertEqual(r.status_code,code,r.text)
        self.assertEqual(self.db.query(CoreCostPriceRevision).count(),2)
        self.assertEqual(self.client.post(self.base,json={'product_code':'B','price':5}).status_code,200)
        self.assertEqual(self.client.post(self.base,json={'product_code':'B','price':0}).status_code,422)
        self.assertEqual(self.client.post(self.base,json={'product_code':'B','price':5,'discount':3.5}).status_code,422)

    def test_legacy_import_blanks_do_not_overwrite_and_brand_comes_from_catalog(self):
        d=self.preview(['货品编号','货品名称','品牌','核心成本价'],[['A','A','错误品牌',None],['B','B','错误品牌',9.5]])
        self.assertEqual(d['errors'],[])
        self.assertEqual(d['blank_price_count'],1)
        self.assertEqual(d['rows'][0]['brand'],'品牌甲')
        r=self.client.post(self.base+'/import',json={'rows':d['rows']})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(router.current_prices(self.db)['A'].price,350)
        self.assertEqual(router.current_prices(self.db)['B'].price,Decimal('9.5'))

    def test_eight_columns_discount_import_rejects_inconsistent_cost(self):
        d=self.preview([label for _,label in router.PRICE_COLUMNS],[['A','A','品牌甲','正装',1000,999,'3.5折',350],['C','C','品牌甲','小样',50,18,5,24]])
        self.assertEqual(len(d['rows']),1)
        self.assertEqual(len(d['errors']),1)
        self.assertEqual(d['rows'][0]['price'],'350.00')
        self.assertEqual(self.client.post(self.base+'/import',json={'rows':d['rows']}).status_code,200)

    def test_export_eight_columns_roundtrip(self):
        r=self.client.get(self.base+'/export')
        self.assertEqual(r.status_code,200,r.text)
        book=load_workbook(BytesIO(r.content),read_only=True)
        self.assertEqual(list(next(book.active.values)),[label for _,label in router.PRICE_COLUMNS])
        book.close()
        d=self.client.post(self.base+'/import-preview',files={'file':('prices.xlsx',r.content)}).json()['data']
        self.assertEqual(d['errors'],[])
        self.assertEqual(len(d['rows']),2)
        self.assertEqual(d['blank_price_count'],3)

    def test_confirm_rechecks_versions_before_any_writes(self):
        d=self.preview(['商品编号','核心成本价'],[['A',300],['C',10]])
        self.client.post(self.base,json={'product_code':'A','price':320})
        self.assertEqual(self.client.post(self.base+'/import',json={'rows':d['rows']}).status_code,409)
        self.assertNotIn('C',router.current_prices(self.db))

    def test_batch_skips_fixed_missing_and_conflicting_rows_and_requires_scope(self):
        r=self.client.post(self.base+'/batch-preview',json={'brand':['品牌甲'],'discount':3.5})
        self.assertEqual(r.status_code,200,r.text)
        d=r.json()['data']
        self.assertEqual((d['matched'],d['skipped'],len(d['errors'])),(4,1,2))
        self.assertEqual([row['product_code'] for row in d['rows']],['C'])
        self.assertNotIn('C',router.current_prices(self.db))
        self.assertEqual(self.client.post(self.base+'/batch',json={'rows':d['rows']}).status_code,200)
        self.assertEqual(router.current_prices(self.db)['C'].price,Decimal('17.5'))
        self.assertEqual(router.current_prices(self.db)['C'].source,'batch')
        self.assertEqual(self.client.post(self.base+'/batch-preview',json={'discount':3.5}).status_code,422)
        tiny=self.client.post(self.base+'/batch-preview',json={'brand':['品牌甲'],'discount':0.0001})
        self.assertEqual(tiny.status_code,200,tiny.text)
        self.assertEqual(tiny.json()['data']['rows'],[])
        self.assertEqual(len(tiny.json()['data']['errors']),3)

    def test_batch_overwrite_requires_option_and_confirms_current_retail(self):
        d=self.client.post(self.base+'/batch-preview',json={'brand':['品牌甲'],'product_type':['正装'],'discount':3.5,'overwrite_fixed':True}).json()['data']
        self.assertEqual([row['product_code'] for row in d['rows']],['A'])
        self.db.execute(text("UPDATE `总库存查询` SET `零售价`=1100 WHERE `货品编号`='A'"));self.db.commit()
        self.assertEqual(self.client.post(self.base+'/batch',json={'rows':d['rows']}).status_code,409)

    def test_delete_restore_preserve_saved_discount_and_history(self):
        self.client.post(self.base,json={'product_code':'A','pricing_mode':'discount','discount':3.5})
        old=router.current_prices(self.db)['A'].id
        self.assertEqual(self.client.delete(self.base+'/A').status_code,200)
        self.assertNotIn('A',router.current_prices(self.db))
        self.assertEqual(self.client.post(self.base+'/A/restore',params={'revision_id':old}).status_code,200)
        saved=router.current_prices(self.db)['A']
        self.assertEqual((saved.price,saved.discount,saved.base_retail_price),(Decimal('350'),Decimal('3.5'),Decimal('1000')))
        history=self.client.get(self.base+'/A/history').json()['data']
        self.assertEqual(history[0]['operator'],'tester')
        self.assertEqual(history[0]['discount'],3.5)

    def test_malformed_file_and_manage_permission(self):
        self.assertEqual(self.client.post(self.base+'/import-preview',files={'file':('bad.xlsx',b'not excel')}).status_code,400)
        self.user.role.permissions=[SimpleNamespace(code='inventory.view')]
        self.assertEqual(self.client.get(self.base).status_code,403)
        self.assertEqual(self.client.post(self.base,json={'product_code':'A','price':1}).status_code,403)

    def test_legacy_schema_migration_is_repeatable_and_keeps_amount(self):
        from app.db import init_db
        engine=create_engine('sqlite://')
        with engine.begin() as c:
            c.execute(text('CREATE TABLE core_cost_price_revisions (id INTEGER PRIMARY KEY, price NUMERIC)'))
            c.execute(text('INSERT INTO core_cost_price_revisions VALUES (1,350)'))
        with patch.object(init_db,'engine',engine):
            init_db.ensure_core_cost_columns();init_db.ensure_core_cost_columns()
        with engine.connect() as c:
            self.assertEqual(tuple(c.execute(text('SELECT price,pricing_mode,discount FROM core_cost_price_revisions')).one()),(350,'fixed',None))
        engine.dispose()


if __name__=='__main__':
    unittest.main()
