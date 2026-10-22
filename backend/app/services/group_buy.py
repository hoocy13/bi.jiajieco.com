"""Monthly group-buy documents; mutations use a batch-wide optimistic revision.

These bounded business documents live in the admin DB, never ODS/ADS. Column IDs
remain stable even when two source/customer labels in an Excel sheet are equal.
"""
from copy import deepcopy
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile, BadZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import text, update
from sqlalchemy.orm import Session
import xlsxwriter

from app.models.group_buy import GroupBuyBatch, GroupBuyChange


def fail(message: str, status: int = 422):
    raise HTTPException(status, detail={"code": status, "message": message, "data": None})


def uid():
    return uuid4().hex


def lookup_product(ods, code):
    """Exact, read-only SKU lookup using the existing BI catalog source priority."""
    code = code.strip()
    if not code:
        fail("请填写产品代码")
    for table in ("总库存查询", "分仓库查询"):
        rows = ods.execute(text(f"""
            SELECT `货品编号` AS code, TRIM(`货品名称`) AS name, TRIM(`条码`) AS barcode,
                   MAX(`updatetime`) AS updated_at
            FROM `{table}` WHERE `货品编号` = :code
            GROUP BY `货品编号`, TRIM(`货品名称`), TRIM(`条码`)
            LIMIT 2
        """), {"code": code}).mappings().all()
        if not rows:
            continue
        if len(rows) > 1:
            return {"status": "ambiguous", "code": code, "source": table,
                    "message": "同一编码存在多份名称或条码，请先核对商品库"}
        row = rows[0]
        if not row["name"] or not row["barcode"]:
            return {"status": "incomplete", "code": code, "source": table,
                    "message": "商品库名称或条码不完整，请核对后填写"}
        return {"status": "matched", "product": {"code": str(row["code"]), "name": str(row["name"]),
                "barcode": str(row["barcode"]), "source": table,
                "updated_at": row["updated_at"].isoformat() if row["updated_at"] else None}}
    return {"status": "not_found", "code": code, "message": "商品库未找到此编码，请核对产品代码"}


def empty_document():
    return {"sources": [], "columns": [], "products": []}


def get_batch(db: Session, batch_id: int, include_deleted=False):
    batch = db.get(GroupBuyBatch, batch_id)
    if batch is None or (batch.status == "deleted" and not include_deleted):
        fail("团购批次不存在", 404)
    return batch


def check_revision(batch, revision, editable=True):
    if batch.revision != revision:
        fail("此批次已被其他人修改，请保留当前输入并刷新后重试", 409)
    if editable and batch.status == "confirmed":
        fail("批次已确认，请先重新开放", 409)


def save_document(db, batch, revision, document, action, changes, operator, status=None, name=None):
    if len(document["products"]) > 5000 or len(document["columns"]) > 200 or len(document["sources"]) > 100:
        fail("单批次最多5000件商品、200个客户列、100个来源，请拆分批次")
    values = {"document": document, "revision": revision + 1,
              "status": status or batch.status, "updated_at": datetime.now()}
    if name is not None:
        values["name"] = name
    result = db.execute(update(GroupBuyBatch).where(
        GroupBuyBatch.id == batch.id, GroupBuyBatch.revision == revision,
    ).values(**values),
        execution_options={"synchronize_session": False})
    if result.rowcount != 1:
        db.rollback()
        fail("此批次已更新，请刷新后重试", 409)
    db.add(GroupBuyChange(batch_id=batch.id, revision=revision + 1, action=action,
                          changes=changes, operator=operator))
    db.commit()
    db.expire_all()
    return get_batch(db, batch.id, include_deleted=True)


def totals(document, product):
    allocated = sum(product["allocations"].values())
    by_source = {s["id"]: 0 for s in document["sources"]}
    for c in document["columns"]:
        by_source[c["source_id"]] += product["allocations"].get(c["id"], 0)
    missing = [s for s in by_source if product["quotas"].get(s) is None and by_source[s] > 0]
    over = [s for s, q in by_source.items() if product["quotas"].get(s) is not None and q > product["quotas"][s]]
    quota = sum(v for v in product["quotas"].values() if v is not None)
    counter_ids = {s["id"] for s in document["sources"] if s["counter"]}
    counter_allocated = sum(q for s, q in by_source.items() if s in counter_ids)
    return {"allocated": allocated, "quota": None if missing else quota,
            "known_quota": quota, "remaining": None if missing else quota - allocated,
            "ordered": sum(product["orders"].values()), "counter_allocated": counter_allocated,
            "missing_sources": missing, "over_sources": over}


def summary(document):
    ts = [totals(document, p) for p in document["products"]]
    return {"product_count": len(ts), "ordered": sum(t["ordered"] for t in ts),
            "allocated": sum(t["allocated"] for t in ts), "known_quota": sum(t["known_quota"] for t in ts),
            "remaining": None if any(t["missing_sources"] for t in ts) else sum(t["remaining"] for t in ts),
            "missing_count": sum(bool(t["missing_sources"]) for t in ts),
            "over_count": sum(bool(t["over_sources"]) for t in ts)}


def batch_data(batch):
    return {"id": batch.id, "month": batch.month, "brand": batch.brand, "name": batch.name,
            "revision": batch.revision, "status": batch.status, "created_by": batch.created_by,
            "updated_at": batch.updated_at.isoformat(), "summary": summary(batch.document)}


def find_product(document, product_id):
    item = next((p for p in document["products"] if p["id"] == product_id), None)
    if item is None:
        fail("商品不存在")
    return item


def ensure_source(document, name, counter=False):
    source = next((s for s in document["sources"] if s["name"] == name), None)
    if source is None:
        source = {"id": uid(), "name": name, "counter": counter}
        document["sources"].append(source)
        for p in document["products"]:
            p["quotas"][source["id"]] = 0
    elif counter:
        source["counter"] = True
    return source


def change_cells(document, cells):
    changes = []
    seen = set()
    for cell in cells:
        key = (cell.product_id, cell.kind, cell.column_id)
        if key in seen:
            fail("同一单元格不能在一次保存中重复提交")
        seen.add(key)
        p = find_product(document, cell.product_id)
        mapping = p["orders"] if cell.kind == "order" else p["allocations"]
        if cell.kind == "order":
            column = next((s for s in document["sources"] if s["id"] == cell.column_id and s["counter"]), None)
            label = column["name"] if column else ""
        else:
            column = next((c for c in document["columns"] if c["id"] == cell.column_id), None)
            source = next((s for s in document["sources"] if column and s["id"] == column["source_id"]), None)
            label = f'{source["name"]} / {column["customer"]}' if source else ""
        if column is None:
            fail("分配列或柜台不存在")
        old = mapping.get(cell.column_id, 0)
        if cell.kind == "order":
            q = p["quotas"].get(cell.column_id)
            p["quotas"][cell.column_id] = cell.quantity + ((q or 0) - old)
            if p["quotas"][cell.column_id] < 0:
                fail("调整下单量后额度为负，请先调整追加额度")
        mapping[cell.column_id] = cell.quantity
        if old != cell.quantity:
            changes.append({"code": p["code"], "column": label, "kind": cell.kind,
                            "before": old, "after": cell.quantity})
    return changes


def remove_item(document, collection, item_id):
    """Remove related values atomically and retain complete removed data in audit."""
    item = next((x for x in document[collection] if x["id"] == item_id), None)
    if item is None:
        fail("商品、客户列或来源不存在", 404)
    changes = {"removed": deepcopy(item)}
    if collection == "columns":
        changes["allocations"] = {p["id"]: p["allocations"].pop(item_id) for p in document["products"] if item_id in p["allocations"]}
    elif collection == "sources":
        removed_columns = [c for c in document["columns"] if c["source_id"] == item_id]
        column_ids = {c["id"] for c in removed_columns}
        changes["columns"] = deepcopy(removed_columns)
        changes["products"] = []
        for p in document["products"]:
            changes["products"].append({"id": p["id"], "code": p["code"],
                "orders": {item_id: p["orders"][item_id]} if item_id in p["orders"] else {},
                "quotas": {item_id: p["quotas"][item_id]} if item_id in p["quotas"] else {},
                "allocations": {k: v for k, v in p["allocations"].items() if k in column_ids}})
            p["orders"].pop(item_id, None)
            p["quotas"].pop(item_id, None)
            p["allocations"] = {k: v for k, v in p["allocations"].items() if k not in column_ids}
        document["columns"] = [c for c in document["columns"] if c["id"] not in column_ids]
    document[collection] = [x for x in document[collection] if x["id"] != item_id]
    return changes


def integer(value, where):
    if value in (None, ""):
        return 0
    try:
        n = Decimal(str(value))
        if not n.is_finite() or n < 0 or n != int(n) or n > 100_000_000:
            raise ValueError
        return int(n)
    except (InvalidOperation, ValueError, OverflowError):
        fail(f"{where}：数量或箱规必须为非负整数，不能忽略负数或非数值记录")


def parse_workbook(content: bytes, kind: str):
    """Read the supplied two-level Excel layout, never trust cached totals/formulas."""
    if len(content) > 12 * 1024 * 1024:
        fail("单个文件不能超过12MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                fail("工作簿解压体积过大，请拆分文件")
        book = load_workbook(BytesIO(content), data_only=False, read_only=True)
    except (BadZipFile, ValueError, OSError, KeyError):
        fail("无法读取文件，请上传有效的xlsx工作簿")
    try:
        sheet = book.active
        if sheet.max_row > 5010 or sheet.max_column > 207:
            fail("工作簿超过单批次行列限制")
        rows = list(sheet.values)
        if len(rows) < 9 or list(rows[7][:4]) != ["国际条形码", "产品代码", "产品名称", "价格"]:
            fail("模板不匹配：第8行须为国际条形码、产品代码、产品名称、价格、ABC、PCB")
        if list(rows[7][4:6]) != ["ABC", "PCB"]:
            fail("第8行第5、6列必须为ABC和PCB（箱规）")
        start = 6 if kind == "order" else 7
        columns = []
        for c in range(start, sheet.max_column):
            source = rows[2][c] if kind == "order" else rows[1][c]
            customer = None if kind == "order" else rows[2][c]
            if not source and not customer:
                if any(r[c] not in (None, "") for r in rows[8:]):
                    fail(f"第{c + 1}列有数量但缺少来源表头")
                continue
            if not source or (kind == "allocation" and not customer):
                fail(f"第{c + 1}列缺少来源或客户名称")
            source, customer = str(source).strip(), str(customer).strip() if customer else None
            if not source or len(source) > 128 or (customer is not None and (not customer or len(customer) > 128)):
                fail("来源或客户名称为空或超过128字")
            columns.append({"source": source, "customer": customer, "excel": c + 1})
        if kind == "order" and len({c["source"] for c in columns}) != len(columns):
            fail("下单表存在重复柜台名称，请先核对")
        products, codes = [], set()
        for r, row in enumerate(rows[8:], 9):
            code = str(row[1] or "").strip()
            if not code:
                if any(v not in (None, "") for v in row):
                    fail(f"第{r}行缺少产品代码")
                continue
            if code in codes:
                fail(f"产品代码重复：{code}，请明确区分后再导入")
            codes.add(code)
            try:
                price = Decimal(str(row[3]))
                if not price.is_finite() or price < 0 or price != price.quantize(Decimal("0.01")) or price >= Decimal('1e14'):
                    raise ValueError
            except (InvalidOperation, ValueError, TypeError):
                fail(f"第{r}行价格无效或超过两位小数")
            if not row[2] or len(str(row[2])) > 300 or len(code) > 128:
                fail(f"第{r}行产品名称为空或字段过长")
            abc = str(row[4] or "").strip()
            if len(abc) > 32:
                fail(f"第{r}行ABC标识过长")
            products.append({"code": code, "name": str(row[2]).strip(), "barcode": str(row[0] or ""),
                             "price": str(price.quantize(Decimal("0.01"))), "abc": abc,
                             "pcb": None if row[5] in (None, "") else integer(row[5], f"第{r}行箱规"),
                             "quantities": [integer(row[c["excel"] - 1], f"第{r}行第{c['excel']}列") for c in columns]})
        if not products or not columns:
            fail("工作簿没有商品或有效数量列")
        ledger = []
        if "分配额度" in book.sheetnames:
            ledger_rows = list(book["分配额度"].values)
            if len(ledger_rows) > 100_001:
                fail("额度台账过大，请拆分批次")
            if not ledger_rows or list(ledger_rows[0][:4]) != ["产品代码", "来源", "下单量", "分配额度"]:
                fail("分配额度工作表格式不匹配")
            seen_ledger = set()
            for r, row in enumerate(ledger_rows[1:], 2):
                if not any(v is not None for v in row):
                    continue
                key = (str(row[0] or "").strip(), str(row[1] or "").strip())
                if not all(key) or key in seen_ledger or key[0] not in codes:
                    fail(f"额度台账第{r}行商品/来源无效或重复")
                seen_ledger.add(key)
                ledger.append({"code": key[0], "source": key[1], "quota": None if row[3] is None else integer(row[3], f"额度台账第{r}行")})
        return {"kind": kind, "columns": columns, "products": products, "ledger": ledger}
    finally:
        book.close()


def merge_import(document, order=None, allocation=None):
    doc = deepcopy(document)
    # Existing quantities are never silently overwritten by importing another version.
    if doc["products"]:
        fail("此批次已有商品，请新建批次导入；已有批次通过在线编辑维护")
    inputs = [x for x in (order, allocation) if x]
    if order and allocation:
        left = {p["code"]: p for p in order["products"]}
        right = {p["code"]: p for p in allocation["products"]}
        if left.keys() != right.keys():
            fail("两张表的产品代码不一致，请核对商品清单")
        for code in left:
            if any(left[code][k] != right[code][k] for k in ("price", "name", "barcode", "abc", "pcb")):
                fail(f"两张表商品属性不一致：{code}")
    for data in inputs:
        for c in data["columns"]:
            ensure_source(doc, c["source"], counter=data["kind"] == "order")
    if allocation:
        for i, c in enumerate(allocation["columns"]):
            source = ensure_source(doc, c["source"])
            doc["columns"].append({"id": uid(), "source_id": source["id"], "customer": c["customer"],
                                   "origin": f"原表第{c['excel']}列", "position": i})
    for item in inputs[0]["products"]:
        doc["products"].append({k: v for k, v in item.items() if k != "quantities"} | {
            "id": uid(), "orders": {}, "allocations": {}, "quotas": {s["id"]: 0 for s in doc["sources"]}})
    by_code = {p["code"]: p for p in doc["products"]}
    for data in inputs:
        for item in data["products"]:
            p = by_code[item["code"]]
            for i, c in enumerate(data["columns"]):
                source = ensure_source(doc, c["source"])
                q = item["quantities"][i]
                if data["kind"] == "order":
                    p["orders"][source["id"]] = q
                    p["quotas"][source["id"]] = q
                else:
                    p["allocations"][doc["columns"][i]["id"]] = q
                    if not source["counter"] and q:
                        p["quotas"][source["id"]] = None
    imported_quotas = {}
    for data in inputs:
        for item in data.get("ledger", []):
            key = (item["code"], item["source"])
            if key in imported_quotas and imported_quotas[key] != item["quota"]:
                fail("两张表的额度台账不一致，请使用同一版本导出文件")
            imported_quotas[key] = item["quota"]
    for (code, name), q in imported_quotas.items():
        source = ensure_source(doc, name)
        by_code[code]["quotas"][source["id"]] = q
    return doc


def export_workbook(document, kind):
    output = BytesIO()
    book = xlsxwriter.Workbook(output, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False})
    sheet = book.add_worksheet("Sheet1")
    fmt = book.add_format({"bg_color": "#F8EFF4", "bold": True, "border": 1})
    price_fmt = book.add_format({"num_format": "#,##0.00"})
    qty_fmt = book.add_format({"num_format": "#,##0"})
    sheet.write_column(0, 0, ["sold to", "ship to", "%", "Counter Name", "TTL VALUE", "TTL NET", "TTL QTY"])
    sheet.write_row(7, 0, ["国际条形码", "产品代码", "产品名称", "价格", "ABC", "PCB"], fmt)
    sheet.set_column(0, 1, 18)
    sheet.set_column(2, 2, 48)
    sheet.set_column(3, 3, 14, price_fmt)
    sheet.set_column(4, 5, 9)
    sources = {s["id"]: s for s in document["sources"]}
    cols = [s for s in document["sources"] if s["counter"]] if kind == "order" else document["columns"]
    start = 6 if kind == "order" else 7
    if kind == "allocation":
        sheet.write(2, 6, "汇总", fmt)
        sheet.write(7, 6, "销售确认量", fmt)
    for i, c in enumerate(cols, start):
        if kind == "order":
            sheet.write(2, i, c["name"], fmt)
        else:
            sheet.write(1, i, sources[c["source_id"]]["name"], fmt)
            sheet.write(2, i, c["customer"], fmt)
        sheet.write(7, i, "销售确认量", fmt)
        sheet.set_column(i, i, 13, qty_fmt)
    for r, p in enumerate(document["products"], 8):
        sheet.write_string(r, 0, p["barcode"])
        sheet.write_string(r, 1, p["code"])
        sheet.write_string(r, 2, p["name"])
        sheet.write_number(r, 3, float(p["price"]), price_fmt)
        sheet.write_string(r, 4, p["abc"])
        if p["pcb"] is not None:
            sheet.write_number(r, 5, p["pcb"])
        qs = []
        for i, c in enumerate(cols, start):
            q = (p["orders"] if kind == "order" else p["allocations"]).get(c["id"], 0)
            qs.append(q)
            sheet.write_number(r, i, q, qty_fmt)
        if kind == "allocation":
            sheet.write_number(r, 6, sum(qs), qty_fmt)
    for i, c in enumerate(cols, start):
        qs = [(p["orders"] if kind == "order" else p["allocations"]).get(c["id"], 0) for p in document["products"]]
        sheet.write_number(6, i, sum(qs), qty_fmt)
        sheet.write_number(4, i, float(sum((Decimal(p["price"]) * q for p, q in zip(document["products"], qs)), Decimal(0))), price_fmt)
    if kind == "allocation":
        sheet.write_number(6, 6, sum(sum(p["allocations"].values()) for p in document["products"]), qty_fmt)
        sheet.write_number(4, 6, float(sum((Decimal(p["price"]) * sum(p["allocations"].values()) for p in document["products"]), Decimal(0))), price_fmt)
    sheet.freeze_panes(8, 6)
    # Preserve manual quotas and their unknown state in a separate, explicit ledger.
    ledger = book.add_worksheet("分配额度")
    ledger.write_row(0, 0, ["产品代码", "来源", "下单量", "分配额度", "已分配", "剩余", "额度状态"], fmt)
    row = 1
    for p in document["products"]:
        for s in document["sources"]:
            q = p["quotas"].get(s["id"])
            a = sum(p["allocations"].get(c["id"], 0) for c in document["columns"] if c["source_id"] == s["id"])
            ledger.write_row(row, 0, [p["code"], s["name"], p["orders"].get(s["id"], 0), q, a, None if q is None else q-a, "待维护" if q is None else "已维护"])
            row += 1
    ledger.set_column(0, 1, 24)
    ledger.set_column(2, 6, 14)
    book.close()
    return output.getvalue()
