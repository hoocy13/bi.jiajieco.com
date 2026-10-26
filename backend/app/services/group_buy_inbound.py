"""Read-only ODS ingestion and explicit receipt-to-order links in the admin DB."""
from datetime import datetime, time, timedelta
from decimal import Decimal, InvalidOperation
import re
import json

from sqlalchemy import inspect, select, text, update, tuple_

from app.models.group_buy import GroupBuyBatch
from app.models.group_buy_inbound import InboundAlias, InboundAudit, InboundBinding, InboundCounter, InboundLine, InboundSync, InboundSupplier, InboundSubstitution, InboundReceiptHold, InboundOrderCollection
from app.schemas.group_buy import BRANDS
from app.services.group_buy import check_revision, fail, get_batch


BRAND_NAMES = {'YSL': ['YSL', 'ysl', '圣罗兰']}
WAREHOUSE = '上海仓库新'


def normalize(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def scope_dates(start, end, brand):
    if brand not in BRANDS:
        fail('请选择团购品牌')
    if end < start or (end - start).days > 92:
        fail('入库日期范围须为1至93天')
    return datetime.combine(start, time.min), datetime.combine(end + timedelta(days=1), time.min)


def load_source(ods, start, end, brand):
    begin, finish = scope_dates(start, end, brand)
    names = BRAND_NAMES.get(brand, [brand])
    placeholders = ','.join(':brand'+str(i) for i in range(len(names)))
    params = {'begin': begin, 'finish': finish, 'warehouse': WAREHOUSE, **{'brand'+str(i): n for i, n in enumerate(names)}}
    rows = ods.execute(text(f'''
        SELECT h.`docId` AS doc_id,d.`recId` AS rec_id,h.`入库单号` AS receipt_number,
          h.`入库时间` AS receipt_time,h.`往来单位` AS supplier,h.`入库仓库` AS warehouse,
          h.`备注` AS remark,h.`关联单号` AS related_number,h.`红冲状态` AS reversal_status,
          d.`货品编号` AS code,d.`货品名称` AS name,d.`条码` AS barcode,
          d.`品牌` AS brand,d.`单位` AS unit,d.`数量` AS quantity,d.`批次` AS lot,
          GREATEST(COALESCE(h.`updatetime`,h.`入库时间`),COALESCE(d.`updatetime`,h.`入库时间`)) AS source_updated_at
        FROM `入库查询` h JOIN `入库查询明细` d ON d.`docId`=h.`docId`
        WHERE h.`入库类型编码`='101' AND h.`入库时间`>=:begin AND h.`入库时间`<:finish
          AND TRIM(h.`入库仓库`)=:warehouse
          AND TRIM(d.`品牌`) IN ({placeholders})
        ORDER BY h.`入库时间`,h.`docId`,d.`recId` LIMIT 20001
    '''), params).mappings().all()
    if len(rows) > 20000:
        fail('此范围超过20,000条入库明细，请缩小入库日期范围；未覆盖已有同步数据')
    return [dict(r) for r in rows]


def snapshot(row, brand):
    result = {k: (v.isoformat() if isinstance(v, datetime) else str(v) if v is not None else '') for k, v in row.items()}
    result['barcode'] = str(row.get('barcode') or '').strip()
    result['brand'] = brand
    result['supplier'] = normalize(row.get('supplier'))
    return result


def source_issue(s):
    if normalize(s.get('warehouse')) != WAREHOUSE:
        return '非上海仓库新，不计入团购到货'
    if s.get('source_missing'):
        return '源明细已变更或不在本次完整采集范围，待复核'
    try:
        n = Decimal(s['quantity'])
        if not n.is_finite() or n != int(n) or abs(n) > 100_000_000:
            return '入库数量不是有效整数'
    except (ValueError, InvalidOperation, KeyError, OverflowError):
        return '入库数量无效'
    # Code 1 is the observed original-document state. Other red/blue states need
    # verified reversal relationships before they can contribute signed totals.
    if str(s.get('reversal_status')) != '1':
        return '红蓝冲状态待复核'
    return ''


def audit(db, action, operator, changes, line_id=None):
    db.add(InboundAudit(action=action, operator=operator, changes=changes, line_id=line_id))


def cas_line(db, line, revision, **values):
    result = db.execute(update(InboundLine).where(InboundLine.id == line.id, InboundLine.revision == revision)
                        .values(**values, revision=revision+1), execution_options={'synchronize_session': False})
    if result.rowcount != 1:
        db.rollback()
        fail('入库记录已更新，请刷新后重试', 409)


def sync_source(db, ods, body, operator):
    begin, finish = scope_dates(body.start_date, body.end_date, body.brand)
    rows = load_source(ods, body.start_date, body.end_date, body.brand)
    existing = db.scalars(select(InboundLine).where(InboundLine.brand == body.brand,
        InboundLine.receipt_time >= begin, InboundLine.receipt_time < finish)).all()
    by_key = {(r.doc_id, r.rec_id): r for r in existing}
    # A source correction may move a receipt across date/brand scopes. Identity
    # is global; updating the new scope must not try to insert it a second time.
    incoming_keys = [(str(r['doc_id']), str(r['rec_id'])) for r in rows]
    for offset in range(0, len(incoming_keys), 400):
        known = db.scalars(select(InboundLine).where(tuple_(InboundLine.doc_id, InboundLine.rec_id)
            .in_(incoming_keys[offset:offset+400]))).all()
        by_key.update({(r.doc_id, r.rec_id): r for r in known})
    seen = set()
    now = datetime.now()
    added = changed = 0
    identity = ('supplier', 'barcode', 'brand', 'warehouse', 'reversal_status')
    for row in rows:
        key = (str(row['doc_id']), str(row['rec_id']))
        if key in seen:
            fail('采集结果存在重复明细身份，本次同步未保存')
        seen.add(key)
        s = snapshot(row, body.brand)
        line = by_key.get(key)
        if line is None:
            db.add(InboundLine(doc_id=key[0], rec_id=key[1], brand=body.brand,
                receipt_time=row['receipt_time'], snapshot=s, synced_at=now))
            added += 1
        elif line.snapshot != s:
            status = line.status
            if any(line.snapshot.get(k) != s.get(k) for k in identity) or source_issue(s) or line.snapshot.get('source_missing'):
                status = 'review' if line.status != 'pending' else 'pending'
            cas_line(db, line, line.revision, snapshot=s, brand=body.brand, status=status, receipt_time=row['receipt_time'], synced_at=now)
            audit(db, 'source_update', operator, {'before': line.snapshot, 'after': s}, line.id)
            changed += 1
        else:
            line.synced_at = now
    for line in existing:
        if (line.doc_id, line.rec_id) not in seen and not line.snapshot.get('source_missing'):
            s = {**line.snapshot, 'source_missing': True}
            cas_line(db, line, line.revision, snapshot=s, status='review', synced_at=now)
            audit(db, 'source_missing', operator, {'before': line.snapshot, 'after': s}, line.id)
            changed += 1
    scope = f'{body.brand}|{body.start_date}|{body.end_date}'
    stamp = db.get(InboundSync, scope)
    if stamp is None:
        stamp = InboundSync(scope=scope)
        db.add(stamp)
    stamp.finished_at, stamp.line_count = now, len(rows)
    stamp.source_updated_at = max((str(r.get('source_updated_at') or '') for r in rows), default=None)
    db.commit()
    db.expire_all()
    return {'new': added, 'changed': changed, 'total': len(rows), 'synced_at': now.isoformat(),
            'source_updated_at': stamp.source_updated_at}


def context(db):
    batches = {b.id: b for b in db.scalars(select(GroupBuyBatch).where(GroupBuyBatch.status != 'deleted')).all()}
    counters = {c.id: c for c in db.scalars(select(InboundCounter)).all()}
    aliases = {(a.brand, a.supplier): a.counter_id for a in db.scalars(select(InboundAlias)).all()}
    bindings = {(b.batch_id, b.source_id): b.counter_id for b in db.scalars(select(InboundBinding)).all()}
    substitutions = {s.line_id: s for s in db.scalars(select(InboundSubstitution)).all()}
    return batches, counters, aliases, bindings, substitutions


def counter_source(batch, counter_id, ctx):
    _, counters, _, bindings = ctx[:4]
    counter = counters.get(counter_id)
    if counter is None:
        return None
    sources = batch.document['sources']
    choices = [s for s in sources if s['counter'] and bindings.get((batch.id, s['id'])) == counter_id]
    if not choices:
        choices = [s for s in sources if s['counter'] and normalize(s['name']) == counter.name
                   and (batch.id, s['id']) not in bindings]
    return choices[0] if len(choices) == 1 else None


def valid_target(line, ctx):
    batch = ctx[0].get(line.batch_id)
    if not batch:
        return None
    source = next((s for s in batch.document['sources'] if s['id'] == line.source_id and s['counter']), None)
    product = next((p for p in batch.document['products'] if p['id'] == line.product_id), None)
    substitution = ctx[4].get(line.id)
    if (not source or not product or batch.brand != line.brand or not product['orders'].get(source['id'], 0)
        or not line.target_barcode or product['barcode'].strip() != line.target_barcode
        or (not substitution and line.snapshot['barcode'] != line.target_barcode) or source_issue(line.snapshot)
        or (substitution and Decimal(substitution.source_quantity) != Decimal(line.snapshot['quantity']))):
        return None
    return batch, source, product


def classify(line, ctx, target_batch=None):
    s = line.snapshot
    row = {**s, 'id': line.id, 'revision': line.revision, 'status': 'pending', 'reason': '',
           'batch_id': line.batch_id, 'source_id': line.source_id, 'product_id': line.product_id,
           'counter': '', 'batch_label': '', 'product': '', 'order_quantity': None,
           'note': line.note, 'operator': line.operator, 'candidates': []}
    batches, counters, aliases, _ = ctx[:4]
    substitution = ctx[4].get(line.id)
    row.update(mode='substitute' if substitution else 'normal',
               credit_quantity=substitution.credit_quantity if substitution else None, source_issue=source_issue(s))
    if line.status == 'excluded':
        row.update(status='excluded', reason='已排除')
        return row
    if line.status == 'confirmed':
        target = valid_target(line, ctx)
        if target:
            b, source, p = target
            row.update(status='confirmed', reason='替代到货' if substitution else '已关联', counter=source['name'], product=p['name'],
                batch_label=f'{b.month} · {b.name}', order_quantity=p['orders'][source['id']])
            return row
        row['reason'] = '原关联商品、柜台或源资料已变化，需解除后重试'
        return row
    issue = source_issue(s)
    if issue:
        row['reason'] = issue
        return row
    if line.status == 'review':
        row['reason'] = '源资料已变化，需解除原处理状态后重试'
        return row
    counter_id = aliases.get((line.brand, s['supplier']))
    if not counter_id:
        row['reason'] = '往来单位未对应柜台'
        return row
    row['counter'] = counters[counter_id].name
    if not s['barcode']:
        row['reason'] = '入库缺少条码'
        return row
    matching = []
    target_reason = '本批次柜台尚未绑定或没有此条码下单商品'
    for batch in batches.values():
        if batch.brand != line.brand:
            continue
        source = counter_source(batch, counter_id, ctx)
        if not source:
            continue
        ordered = [p for p in batch.document['products'] if p['orders'].get(source['id'], 0) > 0]
        products = [p for p in ordered if p.get('barcode', '').strip() == s['barcode']]
        if len(products) > 1:
            if target_batch == batch.id:
                target_reason = '同条码对应多个下单商品'
            continue
        if len(products) == 1:
            p = products[0]
            matching.append({'batch_id': batch.id, 'source_id': source['id'], 'product_id': p['id'],
                'label': f'{batch.month} · {batch.name}', 'counter': source['name'], 'product': p['name'],
                'ordered': p['orders'][source['id']]})
        elif target_batch == batch.id and any(not p.get('barcode', '').strip() for p in ordered):
            target_reason = '条码未匹配，部分下单商品缺少条码'
    row['candidates'] = matching
    choices = [c for c in matching if c['batch_id'] == target_batch] if target_batch else matching
    if len(choices) == 1:
        c = choices[0]
        row.update(status='ready', reason='条码匹配，待确认归属', counter=c['counter'], product=c['product'],
            batch_label=c['label'], order_quantity=c['ordered'])
    else:
        row['reason'] = target_reason if target_batch else ('多个团购批次候选，请选择批次' if matching else '条码未找到唯一的下单商品')
    return row


def options(db):
    ctx = context(db)
    return {'brands': BRANDS, 'batches': [{'id': b.id, 'month': b.month, 'brand': b.brand, 'name': b.name,
        'revision': b.revision, 'sources': [s for s in b.document['sources'] if s['counter']]} for b in ctx[0].values()],
        'aliases': [{'id': a.id, 'brand': a.brand, 'supplier': a.supplier,
            'counter': ctx[1][a.counter_id].name, 'supplier_code': a.supplier_code} for a in db.scalars(select(InboundAlias)).all()],
        'suppliers': [{'code': s.code, 'name': s.name, 'category': s.category, 'alias': s.alias}
            for s in db.scalars(select(InboundSupplier).order_by(InboundSupplier.code)).all()]}


def comparison(db, line_id, batch_id=None, source_id=''):
    line = db.get(InboundLine, line_id)
    if line is None:
        fail('入库明细不存在', 404)
    ctx = context(db)
    row = classify(line, ctx, batch_id)
    result = {'receipt': row, 'source_id': '', 'sources': [], 'orders': [],
              'matching_count': 0, 'mapping_matches': False}
    if not batch_id:
        return result
    batch = get_batch(db, batch_id)
    if batch.brand != line.brand:
        fail('入库品牌与团购批次不一致')
    counter_id = ctx[2].get((line.brand, line.snapshot['supplier']))
    mapped = counter_source(batch, counter_id, ctx) if counter_id else None
    sources = [s for s in batch.document['sources'] if s['counter']]
    barcode = line.snapshot['barcode']
    for source in sources:
        ordered = [p for p in batch.document['products'] if p['orders'].get(source['id'], 0) > 0]
        result['sources'].append({'id': source['id'], 'name': source['name'],
            'matching_count': sum(bool(barcode) and p.get('barcode', '').strip() == barcode for p in ordered)})
    if not source_id:
        source_id = line.source_id if line.status == 'confirmed' and line.batch_id == batch_id else mapped['id'] if mapped else ''
    if not source_id:
        return result
    if not any(s['id'] == source_id for s in sources):
        fail('下单柜台不存在')
    result['source_id'] = source_id
    result['mapping_matches'] = bool(mapped and mapped['id'] == source_id)
    received = progress(db, batch_id)['received']
    for product in batch.document['products']:
        ordered = product['orders'].get(source_id, 0)
        if not ordered:
            continue
        quantity = received.get(f"{product['id']}:{source_id}", 0)
        result['orders'].append({'id': product['id'], 'code': product['code'], 'name': product['name'],
            'barcode': product.get('barcode', ''), 'ordered': ordered, 'received': quantity,
            'waiting': max(ordered-quantity, 0),
            'matches_barcode': bool(barcode) and product.get('barcode', '').strip() == barcode})
    result['orders'].sort(key=lambda p: (not p['matches_barcode'], p['code']))
    result['matching_count'] = sum(p['matches_barcode'] for p in result['orders'])
    return result


def save_mapping(db, body, operator):
    batch = get_batch(db, body.batch_id)
    check_revision(batch, body.revision, editable=False)
    source = next((s for s in batch.document['sources'] if s['id'] == body.source_id and s['counter']), None)
    if not source:
        fail('目标下单柜台不存在')
    supplier, name = normalize(body.supplier), normalize(source['name'])
    catalog = db.get(InboundSupplier, body.supplier_code) if body.supplier_code else None
    if body.supplier_code and (catalog is None or supplier not in (normalize(catalog.name), normalize(catalog.alias))):
        fail('供应商编号与所选名称不一致，请从供应商档案重新选择')
    counter = db.scalar(select(InboundCounter).where(InboundCounter.brand == batch.brand, InboundCounter.name == name))
    if not counter:
        counter = InboundCounter(brand=batch.brand, name=name)
        db.add(counter)
        db.flush()
    alias = db.scalar(select(InboundAlias).where(InboundAlias.brand == batch.brand, InboundAlias.supplier == supplier))
    before = alias.counter_id if alias else None
    if alias is None:
        alias = InboundAlias(brand=batch.brand, supplier=supplier, counter_id=counter.id)
        db.add(alias)
    alias.counter_id = counter.id
    alias.supplier_code = body.supplier_code
    binding = db.scalar(select(InboundBinding).where(InboundBinding.batch_id == batch.id, InboundBinding.source_id == source['id']))
    if binding is None:
        db.add(InboundBinding(batch_id=batch.id, source_id=source['id'], counter_id=counter.id))
    else:
        binding.counter_id = counter.id
    audit(db, 'counter_mapping', operator, {'brand': batch.brand, 'supplier': supplier,
        'before_counter_id': before, 'counter_id': counter.id, 'batch_id': batch.id, 'source_id': source['id']})
    db.commit()


def act(db, body, operator):
    if len({x.id for x in body.lines}) != len(body.lines):
        fail('同一明细不能重复提交')
    if body.action in ('exclude', 'release') and not body.note:
        fail('请填写处理原因')
    ctx = context(db)
    if body.action == 'confirm':
        batch = get_batch(db, body.batch_id) if body.batch_id else None
        if not batch or not body.batch_revision:
            fail('请选择并核对目标团购批次')
        check_revision(batch, body.batch_revision, editable=False)
        # Serialize confirmation against concurrent order/product deletion without
        # modifying quantities or advancing the order-document revision.
        locked = db.execute(update(GroupBuyBatch).where(GroupBuyBatch.id == batch.id,
            GroupBuyBatch.revision == body.batch_revision).values(revision=body.batch_revision),
            execution_options={'synchronize_session': False})
        if locked.rowcount != 1:
            db.rollback()
            fail('团购批次已更新，请刷新后重试', 409)
    for version in body.lines:
        line = db.get(InboundLine, version.id)
        if not line or line.revision != version.revision:
            db.rollback()
            fail('入库明细已被更新，请刷新后重试', 409)
        before = {'status': line.status, 'batch_id': line.batch_id, 'source_id': line.source_id, 'product_id': line.product_id}
        if body.action == 'confirm':
            ensure_receipt_available(db, line)
            row = classify(line, ctx, body.batch_id)
            if row['status'] != 'ready':
                db.rollback()
                fail('所选明细存在不能确认的项目：'+row['reason'], 409)
            c = next(c for c in row['candidates'] if c['batch_id'] == body.batch_id)
            values = {'status': 'confirmed', 'batch_id': body.batch_id, 'source_id': c['source_id'],
                'product_id': c['product_id'], 'target_barcode': line.snapshot['barcode']}
        elif body.action == 'exclude':
            if line.status == 'confirmed':
                db.rollback()
                fail('已关联明细请先解除关联', 409)
            values = {'status': 'excluded', 'batch_id': None, 'source_id': None, 'product_id': None, 'target_barcode': None}
        else:
            values = {'status': 'pending', 'batch_id': None, 'source_id': None, 'product_id': None, 'target_barcode': None}
        substitution = db.get(InboundSubstitution, line.id)
        if substitution:
            before['substitution'] = {'credit_quantity': substitution.credit_quantity, 'source_quantity': substitution.source_quantity}
            db.delete(substitution)
        cas_line(db, line, version.revision, **values, note=body.note, operator=operator)
        audit(db, body.action, operator, {'before': before, 'after': values, 'note': body.note}, line.id)
    db.commit()
    db.expire_all()
    return {'processed': len(body.lines)}


def inbound_product_key(row):
    return json.dumps([row.get('code', ''), row.get('barcode', '')], ensure_ascii=False, separators=(',', ':'))


def list_lines(db, start, end, brand, batch_id, state, keyword, supplier, page, page_size, sort_by, sort_order,
               supplier_names=None, product_keys=None):
    begin, finish = scope_dates(start, end, brand)
    query = select(InboundLine).where(InboundLine.brand == brand, InboundLine.receipt_time >= begin, InboundLine.receipt_time < finish)
    lines = db.scalars(query).all()
    ctx = context(db)
    rows = [classify(line, ctx, batch_id) for line in lines if normalize(line.snapshot.get('warehouse')) == WAREHOUSE]
    suppliers = sorted({r['supplier'] for r in rows})
    if supplier:
        rows = [r for r in rows if r['supplier'] == normalize(supplier)]
    if keyword:
        q = keyword.strip().lower()
        rows = [r for r in rows if q in ' '.join(str(r.get(k, '')) for k in ('receipt_number', 'barcode', 'code', 'name')).lower()]
    if batch_id:
        rows = [r for r in rows if r['status'] != 'confirmed' or r['batch_id'] == batch_id]
    # Column choices span the full scope, independently of page and selections.
    supplier_filters = [{'value': name, 'label': name, 'search': name} for name in sorted({r['supplier'] for r in rows})]
    products = {}
    for row in rows:
        product_key = inbound_product_key(row)
        products[product_key] = {'value': product_key, 'label': row['name'],
                                'detail': ' · '.join(filter(None, (row['code'], row['barcode']))),
                                'search': ' '.join((row['name'], row['code'], row['barcode']))}
    product_filters = sorted(products.values(), key=lambda p: (p['label'], p['value']))
    if supplier_names:
        names = {normalize(name) for name in supplier_names}
        rows = [r for r in rows if r['supplier'] in names]
    if product_keys:
        keys = set(product_keys)
        rows = [r for r in rows if inbound_product_key(r) in keys]
    summary = {s: sum(r['status'] == s for r in rows) for s in ('pending', 'ready', 'confirmed', 'excluded')}
    if state == 'pending':
        rows = [r for r in rows if r['status'] in ('pending', 'ready')]
    elif state != 'all':
        rows = [r for r in rows if r['status'] == state]
    if batch_id and state == 'confirmed':
        rows = [r for r in rows if r['batch_id'] == batch_id]
    def sort_key(row):
        value = row.get(sort_by, '')
        if sort_by == 'quantity':
            try:
                value = Decimal(value)
                if not value.is_finite():
                    value = Decimal(0)
            except (InvalidOperation, ValueError):
                value = Decimal(0)
        return value, row['id']
    rows.sort(key=sort_key, reverse=sort_order == 'desc')
    total = len(rows)
    stamp = db.get(InboundSync, f'{brand}|{start}|{end}')
    return {'rows': rows[(page-1)*page_size:page*page_size], 'total': total, 'summary': summary,
        'suppliers': suppliers, 'receipt_count': len({r['doc_id'] for r in rows}),
        'supplier_filters': supplier_filters, 'product_filters': product_filters,
        'synced_at': stamp.finished_at.isoformat() if stamp else None,
        'source_updated_at': stamp.source_updated_at if stamp else None}


def progress(db, batch_id, product_id='', source_id=''):
    batch = get_batch(db, batch_id)
    ctx = context(db)
    lines = db.scalars(select(InboundLine).where(InboundLine.batch_id == batch.id, InboundLine.status == 'confirmed')).all()
    received, normal_received, substituted, details, review = {}, {}, {}, [], 0
    for line in lines:
        if not valid_target(line, ctx):
            review += 1
            continue
        key = line.product_id+':'+line.source_id
        substitution = ctx[4].get(line.id)
        quantity = substitution.credit_quantity if substitution else int(Decimal(line.snapshot['quantity']))
        received[key] = received.get(key, 0) + quantity
        bucket = substituted if substitution else normal_received
        bucket[key] = bucket.get(key, 0) + quantity
        if (product_id or source_id) and (not product_id or product_id == line.product_id) and (not source_id or source_id == line.source_id):
            details.append(classify(line, ctx, batch.id))
    review += db.scalar(select(text('count(*)')).select_from(InboundLine).where(InboundLine.batch_id == batch.id, InboundLine.status == 'review')) or 0
    stamp = db.scalar(select(InboundSync).where(InboundSync.scope.like(batch.brand+'|%')).order_by(InboundSync.finished_at.desc()).limit(1))
    products = {}
    summary = {'ordered': 0, 'received_total': 0, 'waiting_total': 0, 'over_received': 0}
    for product in batch.document['products']:
        totals = dict.fromkeys(summary, 0)
        for source in batch.document['sources']:
            if not source['counter']:
                continue
            ordered = product['orders'].get(source['id'], 0)
            actual = received.get(product['id']+':'+source['id'], 0)
            totals['ordered'] += ordered
            totals['received_total'] += actual
            totals['waiting_total'] += max(ordered-actual, 0)
            totals['over_received'] += max(actual-ordered, 0)
        products[product['id']] = totals
        for name in summary:
            summary[name] += totals[name]
    return {'received': received, 'normal_received': normal_received, 'substituted': substituted, 'details': details, 'review_count': review,
        'products': products, 'summary': summary,
        'synced_at': stamp.finished_at.isoformat() if stamp else None,
        'source_updated_at': stamp.source_updated_at if stamp else None}


def hold_data(hold):
    return {'revision': hold.revision if hold else 0, 'waiting': bool(hold and hold.waiting),
            'month': hold.month if hold else '', 'counter': hold.counter if hold else '',
            'note': hold.note if hold else ''}


def ensure_receipt_available(db, line):
    hold = db.get(InboundReceiptHold, (line.doc_id, line.brand))
    if hold and hold.waiting:
        db.rollback()
        fail('该单待补下单表，请先选择“资料齐了，重新匹配”', 409)
    # Serialize confirmation against an existing hold's concurrent status update.
    if hold:
        locked = db.execute(update(InboundReceiptHold).where(InboundReceiptHold.doc_id == line.doc_id,
            InboundReceiptHold.brand == line.brand, InboundReceiptHold.revision == hold.revision,
            InboundReceiptHold.waiting.is_(False)).values(revision=hold.revision), execution_options={'synchronize_session': False})
        if locked.rowcount != 1:
            db.rollback()
            fail('入库单待补状态已变化，请刷新', 409)


def set_receipt_hold(db, doc_id, brand, body, operator):
    try:
        detail = receipt_detail(db, doc_id, brand)
        if body.waiting and not any(r['status'] in ('pending', 'ready') for r in detail['rows']):
            fail('该单已处理完，无需等待下单表')
        hold = db.get(InboundReceiptHold, (doc_id, brand))
        before = hold_data(hold)
        if body.revision != before['revision']:
            fail('待补状态已更新，请刷新后重试', 409)
        values = dict(waiting=body.waiting, month=body.month, counter=normalize(body.counter),
                      note=body.note, operator=operator, updated_at=datetime.now())
        if hold:
            result = db.execute(update(InboundReceiptHold).where(InboundReceiptHold.doc_id == doc_id,
                InboundReceiptHold.brand == brand, InboundReceiptHold.revision == body.revision)
                .values(**values, revision=body.revision+1), execution_options={'synchronize_session': False})
            if result.rowcount != 1:
                fail('待补状态已更新，请刷新后重试', 409)
        else:
            db.add(InboundReceiptHold(doc_id=doc_id, brand=brand, revision=1, **values))
        if body.waiting:
            record = db.scalar(select(InboundOrderCollection).where(InboundOrderCollection.month==body.month,
                InboundOrderCollection.brand==brand, InboundOrderCollection.counter==normalize(body.counter),
                InboundOrderCollection.batch_key==''))
            if not record:
                db.add(InboundOrderCollection(month=body.month,brand=brand,counter=normalize(body.counter),
                    batch_key='',status='missing',note=body.note,operator=operator))
        audit(db, 'receipt_waiting' if body.waiting else 'receipt_resume', operator,
              {'doc_id':doc_id,'brand':brand,'before':before,'after':{**values,'updated_at':values['updated_at'].isoformat()},'note':body.note})
        db.commit();db.expire_all()
        return hold_data(db.get(InboundReceiptHold,(doc_id,brand)))
    except Exception:
        db.rollback();raise


def list_order_collections(db, month, brand):
    if brand not in BRANDS:
        fail('请选择团购品牌')
    batches = context(db)[0]
    records = db.scalars(select(InboundOrderCollection).where(InboundOrderCollection.month==month,
        InboundOrderCollection.brand==brand).order_by(InboundOrderCollection.counter,InboundOrderCollection.id)).all()
    return [{'id':r.id,'month':r.month,'brand':r.brand,'counter':r.counter,'batch_id':int(r.batch_key) if r.batch_key else None,
        'batch_name':batches[int(r.batch_key)].name if r.batch_key and int(r.batch_key) in batches else '批次待定',
        'status':r.status,'note':r.note,'revision':r.revision} for r in records]


def save_order_collection(db, body, operator):
    if body.brand not in BRANDS:
        fail('请选择团购品牌')
    if body.batch_id:
        batch = get_batch(db,body.batch_id)
        if batch.brand!=body.brand or batch.month!=body.month:
            fail('批次与月份、品牌不一致')
    try:
        record = db.scalar(select(InboundOrderCollection).where(InboundOrderCollection.month==body.month,
            InboundOrderCollection.brand==body.brand,InboundOrderCollection.counter==normalize(body.counter),
            InboundOrderCollection.batch_key==str(body.batch_id or '')))
        if body.revision != (record.revision if record else 0):
            fail('收集状态已更新，请刷新',409)
        before = {'status':record.status,'note':record.note} if record else None
        if record:
            result=db.execute(update(InboundOrderCollection).where(InboundOrderCollection.id==record.id,
                InboundOrderCollection.revision==body.revision).values(status=body.status,note=body.note,
                operator=operator,updated_at=datetime.now(),revision=body.revision+1),execution_options={'synchronize_session':False})
            if result.rowcount!=1:
                fail('收集状态已更新，请刷新',409)
        else:
            db.add(InboundOrderCollection(month=body.month,brand=body.brand,counter=normalize(body.counter),
                batch_key=str(body.batch_id or ''),status=body.status,note=body.note,operator=operator))
        audit(db,'order_collection',operator,{'before':before,'after':body.model_dump()})
        db.commit();db.expire_all()
        return list_order_collections(db,body.month,body.brand)
    except Exception:
        db.rollback();raise


def recommend_receipt(db, lines, rows, ctx):
    """Read-only recommendations; never resolve competing batches by date alone."""
    received = {}
    for line in db.scalars(select(InboundLine).where(InboundLine.status == 'confirmed', InboundLine.brand == lines[0].brand)).all():
        if valid_target(line, ctx):
            key = (line.batch_id, line.source_id, line.product_id)
            sub = ctx[4].get(line.id)
            received[key] = received.get(key, 0) + (sub.credit_quantity if sub else int(Decimal(line.snapshot['quantity'])))
    reservations, scores = {}, {}
    for line, row in zip(lines, rows):
        row['recommendation'] = None
        row['auto_match_reason'] = row['reason']
        if line.status != 'pending' or source_issue(line.snapshot):
            continue
        candidates = []
        duplicate = False
        counter_id = ctx[2].get((line.brand, line.snapshot['supplier']))
        for b in ctx[0].values():
            if b.brand != line.brand or not counter_id:
                continue
            source = counter_source(b, counter_id, ctx)
            if source and line.snapshot['barcode']:
                duplicate |= sum(p['orders'].get(source['id'], 0)>0 and p.get('barcode', '').strip()==line.snapshot['barcode']
                                 for p in b.document['products']) > 1
        for c in row['candidates']:
            key = (c['batch_id'], c['source_id'], c['product_id'])
            candidates.append({**c, 'waiting': max(c['ordered']-received.get(key, 0), 0)})
            scores[c['batch_id']] = scores.get(c['batch_id'], 0)+1
        if duplicate:
            row['auto_match_reason'] = '同条码存在重复下单商品，请人工核对'
            continue
        quantity = int(Decimal(line.snapshot['quantity']))
        fitting = [c for c in candidates if quantity>0 and c['waiting']>=quantity]
        chosen = fitting[0] if len(fitting)==1 else candidates[0] if len(candidates)==1 else None
        if chosen:
            automatic = len(fitting)==1
            row['recommendation'] = {**chosen, 'automatic': automatic}
            row['auto_match_reason'] = '自动匹配，待确认' if automatic else '已推荐唯一归属，数量超过待到货或非正数，请核对'
            if automatic:
                key = (chosen['batch_id'], chosen['source_id'], chosen['product_id'])
                reservations.setdefault(key, []).append((row, quantity))
        elif candidates:
            row['auto_match_reason'] = '多个批次均可匹配，请人工选择' if fitting else '候选批次待到货不足，请人工核对'
    for group in reservations.values():
        if sum(q for _, q in group)>group[0][0]['recommendation']['waiting']:
            for row, _ in group:
                row['recommendation']['automatic'] = False
                row['auto_match_reason'] = '同单同商品合计超过待到货，请核对归属'
    ranked = sorted(scores, key=lambda k: (-scores[k], k))
    return ranked[0] if ranked and (len(ranked)==1 or scores[ranked[0]]>scores[ranked[1]]) else None


def receipt_detail(db, doc_id, brand, batch_id=None):
    if brand not in BRANDS:
        fail('请选择团购品牌')
    lines = db.scalars(select(InboundLine).where(InboundLine.doc_id == doc_id, InboundLine.brand == brand)
                       .order_by(InboundLine.id)).all()
    lines = [line for line in lines if normalize(line.snapshot.get('warehouse')) == WAREHOUSE]
    if not lines:
        fail('上海仓库新未找到该入库单', 404)
    ctx = context(db)
    rows = [classify(line, ctx, batch_id) for line in lines]
    recommended_batch = recommend_receipt(db, lines, rows, ctx)
    hold = db.get(InboundReceiptHold, (doc_id, brand))
    waiting = bool(hold and hold.waiting)
    if waiting:
        for row in rows:
            if row['recommendation']:
                row['recommendation']['automatic'] = False
            row['auto_match_reason'] = '待补下单表，资料齐后重新匹配'
    return {'doc_id': doc_id, 'receipt_number': lines[0].snapshot['receipt_number'],
            'supplier': lines[0].snapshot['supplier'], 'warehouse': WAREHOUSE,
            'receipt_time': lines[0].receipt_time.isoformat(),
            'rows': rows, 'recommended_batch_id': recommended_batch, 'hold': hold_data(hold),
            'auto_count': sum(bool(r['recommendation'] and r['recommendation']['automatic']) for r in rows)}


def list_receipts(db, start, end, brand, state='pending', keyword='', supplier='', page=1, page_size=30,
                  sort_by='receipt_time', sort_order='desc', supplier_names=None, product_keys=None):
    begin, finish = scope_dates(start, end, brand)
    lines = db.scalars(select(InboundLine).where(InboundLine.brand == brand,
        InboundLine.receipt_time >= begin, InboundLine.receipt_time < finish)).all()
    ctx = context(db)
    rows = [classify(line, ctx) for line in lines if normalize(line.snapshot.get('warehouse')) == WAREHOUSE]
    suppliers = sorted({r['supplier'] for r in rows})
    products = {inbound_product_key(r): {'value': inbound_product_key(r), 'label': r['name'],
        'detail': r['code']+' · '+r['barcode'], 'search': r['name']+' '+r['code']+' '+r['barcode']} for r in rows}
    groups = {}
    for row in rows:
        groups.setdefault(row['doc_id'], []).append(row)
    documents = []
    holds = {h.doc_id: h for h in db.scalars(select(InboundReceiptHold).where(InboundReceiptHold.brand == brand)).all()}
    for doc_id, items in groups.items():
        first = items[0]
        if supplier and normalize(supplier) != first['supplier']:
            continue
        if supplier_names and first['supplier'] not in {normalize(n) for n in supplier_names}:
            continue
        if product_keys and not any(inbound_product_key(r) in product_keys for r in items):
            continue
        if keyword and not any(keyword.strip().lower() in ' '.join(str(r.get(k, '')) for k in
                ('receipt_number', 'barcode', 'code', 'name')).lower() for r in items):
            continue
        completed = sum(r['status'] in ('confirmed', 'excluded') for r in items)
        excluded = sum(r['status'] == 'excluded' for r in items)
        status = 'excluded' if excluded == len(items) else 'done' if completed == len(items) else 'partial' if completed else 'pending'
        hold = holds.get(doc_id)
        if completed < len(items) and hold and hold.waiting:
            status = 'waiting'
        documents.append({'doc_id': doc_id, 'receipt_number': first['receipt_number'], 'supplier': first['supplier'],
            'receipt_time': first['receipt_time'], 'warehouse': WAREHOUSE, 'item_count': len(items),
            'processed_count': completed, 'pending_count': len(items)-completed, 'status': status,
            'confirmed_count': completed-excluded, 'excluded_count': excluded, 'hold': hold_data(hold),
            'destinations': sorted({r['batch_label']+' · '+r['counter'] for r in items if r['status']=='confirmed'}),
            'lines': [{'id': r['id'], 'revision': r['revision'], 'status': r['status']} for r in items]})
    summary = {s: sum(d['status'] == s for d in documents) for s in ('pending', 'partial', 'waiting', 'done', 'excluded')}
    if state == 'pending':
        documents = [d for d in documents if d['status'] in ('pending', 'partial')]
    elif state != 'all':
        documents = [d for d in documents if d['status'] == state]
    documents.sort(key=lambda d: (d.get(sort_by, d['receipt_time']), d['doc_id']), reverse=sort_order=='desc')
    stamp = db.get(InboundSync, f'{brand}|{start}|{end}')
    return {'rows': documents[(page-1)*page_size:page*page_size], 'total': len(documents), 'summary': summary,
        'suppliers': suppliers, 'supplier_filters': [{'value': s, 'label': s} for s in suppliers],
        'product_filters': sorted(products.values(), key=lambda p: p['label']),
        'synced_at': stamp.finished_at.isoformat() if stamp else None,
        'source_updated_at': stamp.source_updated_at if stamp else None}


def confirm_receipt(db, body, operator):
    """Validate all selected lines, lock batches in ID order, then commit atomically."""
    if len({v.id for v in body.lines}) != len(body.lines):
        fail('同一明细不能重复提交')
    try:
        ctx = context(db)
        revisions = {}
        for v in body.lines:
            if v.batch_id in revisions and revisions[v.batch_id] != v.batch_revision:
                fail('同一批次版本不一致，请刷新', 409)
            revisions[v.batch_id] = v.batch_revision
        for batch_id, revision in sorted(revisions.items()):
            batch = get_batch(db, batch_id)
            check_revision(batch, revision, editable=False)
            locked = db.execute(update(GroupBuyBatch).where(GroupBuyBatch.id == batch_id,
                GroupBuyBatch.revision == revision).values(revision=revision), execution_options={'synchronize_session': False})
            if locked.rowcount != 1:
                fail('团购批次已更新，请刷新后重试', 409)
        changes = []
        for v in body.lines:
            line = db.get(InboundLine, v.id)
            if not line or line.revision != v.revision:
                fail('入库明细已更新，请刷新后重试', 409)
            ensure_receipt_available(db, line)
            if line.status != 'pending' or source_issue(line.snapshot):
                fail('所选明细不能确认：'+(source_issue(line.snapshot) or '请先解除原处理状态'), 409)
            batch = ctx[0].get(v.batch_id)
            if not batch or batch.brand != line.brand:
                fail('入库品牌与团购批次不一致')
            source = next((s for s in batch.document['sources'] if s['id']==v.source_id and s['counter']), None)
            if not source:
                fail('所选柜台不存在')
            ordered = [p for p in batch.document['products'] if p['orders'].get(source['id'], 0)>0]
            if v.mode == 'normal':
                matches = [p for p in ordered if line.snapshot['barcode'] and p.get('barcode', '').strip()==line.snapshot['barcode']]
                if len(matches) != 1:
                    fail('正常到货须唯一匹配条码；替代补货请选择“替代到货”')
                product = matches[0]
                if v.product_id and v.product_id != product['id']:
                    fail('所选下单商品与条码不一致')
            else:
                product = next((p for p in ordered if p['id']==v.product_id), None)
                if not product or not product.get('barcode', '').strip():
                    fail('请选择有条码的被替代下单商品')
                if not v.credit_quantity or not v.note or Decimal(line.snapshot['quantity'])<=0:
                    fail('替代到货需正数入库数量、抵扣数量和替代原因')
                if product.get('barcode', '').strip() == line.snapshot['barcode']:
                    fail('条码相同请使用正常到货')
            values = {'status': 'confirmed', 'batch_id': batch.id, 'source_id': source['id'],
                'product_id': product['id'], 'target_barcode': product.get('barcode', '').strip()}
            changes.append((line, v, values))
        for line, v, values in changes:
            before = {'status': line.status, 'batch_id': line.batch_id, 'source_id': line.source_id, 'product_id': line.product_id}
            cas_line(db, line, v.revision, **values, note=v.note, operator=operator)
            if v.mode == 'substitute':
                db.add(InboundSubstitution(line_id=line.id, credit_quantity=v.credit_quantity, source_quantity=line.snapshot['quantity']))
            audit(db, 'confirm', operator, {'before': before, 'after': values, 'mode': v.mode,
                'source_quantity': line.snapshot['quantity'], 'credit_quantity': v.credit_quantity,
                'note': v.note, 'supplier': line.snapshot['supplier']}, line.id)
        db.commit()
        db.expire_all()
        return {'processed': len(changes)}
    except Exception:
        db.rollback()
        raise


def guard_order_structure(db, batch_id, product_id=None, source_id=None):
    if not inspect(db.get_bind()).has_table(InboundLine.__tablename__):
        return
    query = select(InboundLine.id).where(InboundLine.batch_id == batch_id,
        InboundLine.status.in_(('confirmed', 'review')))
    if product_id:
        query = query.where(InboundLine.product_id == product_id)
    if source_id:
        query = query.where(InboundLine.source_id == source_id)
    if db.scalar(query.limit(1)) is not None:
        fail('此商品或柜台已有入库关联，请先解除关联再删除或修改条码', 409)


def import_suppliers(db, content, operator):
    from io import BytesIO
    from zipfile import ZipFile, BadZipFile
    from openpyxl import load_workbook
    if len(content) > 5 * 1024 * 1024:
        fail('供应商文件不能超过5MB')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 40 * 1024 * 1024:
                fail('供应商工作簿解压体积过大')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, ValueError, OSError, KeyError):
        fail('请上传有效的供应商xlsx文件')
    parsed = []
    try:
        for sheet in book:
            # Jike exports declare A1:A1 even when the XML contains hundreds of rows.
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, ())
            if not {'编号', '名称'}.issubset(headers):
                continue
            index = {h: i for i, h in enumerate(headers) if h}
            for number, row in enumerate(rows, 2):
                if number > 10002:
                    fail('供应商最多10,000条，请拆分文件')
                if not any(v is not None for v in row):
                    continue
                values = {k: str(row[index[k]] or '').strip() if k in index and index[k] < len(row) else ''
                          for k in ('编号', '名称', '分类', '别名')}
                if not values['编号'] or not values['名称'] or any(v.startswith('=') for v in values.values()):
                    fail(f'供应商第{number}行编号、名称缺失或含公式')
                if len(values['编号']) > 128 or len(values['名称']) > 255 or len(values['分类']) > 128 or len(values['别名']) > 255:
                    fail(f'供应商第{number}行字段过长')
                parsed.append(values)
            break
    finally:
        book.close()
    if not parsed or len({r['编号'] for r in parsed}) != len(parsed):
        fail('供应商清单为空、模板不匹配或编号重复')
    for r in parsed:
        supplier = db.get(InboundSupplier, r['编号'])
        if supplier is None:
            supplier = InboundSupplier(code=r['编号'])
            db.add(supplier)
        supplier.name, supplier.category, supplier.alias = r['名称'], r['分类'], r['别名']
        supplier.imported_at = datetime.now()
    audit(db, 'supplier_import', operator, {'count': len(parsed), 'fields': ['编号', '名称', '分类', '别名']})
    db.commit()
    return {'imported': len(parsed)}
