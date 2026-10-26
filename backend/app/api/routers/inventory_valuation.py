from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from urllib.parse import quote
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile
from openpyxl import load_workbook
import xlsxwriter
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, require_permission
from app.db.ods import get_ods_db
from app.db.session import get_db
from app.models.core_cost_price import CoreCostPriceRevision
from app.models.user import User
from app.schemas.common import ok
from app.schemas.core_cost import PriceInput, ImportInput, BatchPriceInput
from app.services.core_cost_catalog import load_catalog, resolve_price, calculate_cost


router = APIRouter(prefix="/inventory/valuation", tags=["inventory-valuation"])


def price_states(db: Session) -> tuple[dict[str, CoreCostPriceRevision], dict[str, CoreCostPriceRevision]]:
    revisions = db.query(CoreCostPriceRevision).filter(
        CoreCostPriceRevision.effective_date <= date.today()
    ).order_by(
        CoreCostPriceRevision.product_code,
        CoreCostPriceRevision.effective_date.desc(),
        CoreCostPriceRevision.id.desc(),
    ).all()
    deleted_ids = {}
    deleted = {}
    for revision in revisions:
        if revision.source == "delete" and revision.id > deleted_ids.get(revision.product_code, 0):
            deleted_ids[revision.product_code] = revision.id
            deleted[revision.product_code] = revision
    active = {}
    for revision in revisions:
        if revision.source != "delete" and revision.id > deleted_ids.get(revision.product_code, 0):
            active.setdefault(revision.product_code, revision)
    return active, {code: revision for code, revision in deleted.items() if code not in active}


def current_prices(db: Session) -> dict[str, CoreCostPriceRevision]:
    return price_states(db)[0]


def price_data(revision: CoreCostPriceRevision) -> dict:
    return {
        "id": revision.id,
        "product_code": revision.product_code,
        "barcode": revision.barcode,
        "product_name": revision.product_name,
        "price": float(revision.price),
        "pricing_mode": getattr(revision, "pricing_mode", None) or "fixed",
        "discount": float(revision.discount) if getattr(revision, "discount", None) is not None else None,
        "base_retail_price": float(revision.base_retail_price) if getattr(revision, "base_retail_price", None) is not None else None,
        "effective_date": revision.effective_date.isoformat(),
        "source": revision.source,
        "restored_from_id": revision.restored_from_id,
        "operator": revision.operator,
        "created_at": revision.created_at.isoformat() if revision.created_at else None,
    }


def add_revision(db: Session, item: PriceInput, operator: str, source: str, restored_from_id: int | None = None) -> None:
    db.add(CoreCostPriceRevision(
        product_code=item.product_code.strip(), product_name=item.product_name.strip(), barcode=item.barcode,
        price=item.price, effective_date=item.effective_date, source=source,
        pricing_mode=item.pricing_mode, discount=item.discount, base_retail_price=item.base_retail_price,
        restored_from_id=restored_from_id, operator=operator,
    ))


@router.get("")
def valuation(
    keyword: str = "",
    warehouse: list[str] | None = Query(None),
    product_type: list[str] | None = Query(None),
    brand: list[str] | None = Query(None),
    scope: str = Query("total", pattern="^(total|split)$"),
    missing_price: str = Query("all", pattern="^(all|retail|tax|core)$"),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=10, le=100),
    ods: Session = Depends(get_ods_db),
    admin: Session = Depends(get_db),
) -> dict:
    filters = []
    params: dict = {}
    if keyword.strip():
        filters.append("(`货品编号` LIKE :keyword OR `货品名称` LIKE :keyword OR `条码` LIKE :keyword)")
        params["keyword"] = f"%{keyword.strip()}%"
    brands = sorted({value.strip() for value in brand or [] if value.strip()})
    if brands:
        keys = []
        for index, value in enumerate(brands):
            param = f"brand_{index}"
            params[param] = value
            keys.append(f":{param}")
        filters.append(f"`品牌` IN ({', '.join(keys)})")
    for field, values, key in (("仓库", warehouse if scope == "split" else None, "warehouse"), ("货品分类", product_type, "type")):
        values = sorted({value.strip() for value in values or [] if value.strip() and value != "__all__"})
        if values:
            keys = []
            for index, value in enumerate(values):
                param = f"{key}_{index}"
                params[param] = value
                keys.append(f":{param}")
            filters.append(f"`{field}` IN ({', '.join(keys)})")
    where = " AND ".join(filters) if filters else "1 = 1"
    source_table = "总库存查询" if scope == "total" else "分仓库查询"
    warehouse_column = "MAX(`仓库`)" if scope == "total" else "`仓库`"
    group_columns = "`货品编号`, `零售价`, `含税价`" if scope == "total" else "`货品编号`, `仓库`, `零售价`, `含税价`"
    stock_rows = ods.execute(text(f"""
        SELECT `货品编号` AS product_code, MAX(`货品名称`) AS product_name,
          MAX(`品牌`) AS brand, {warehouse_column} AS warehouse, MAX(`货品分类`) AS product_type,
          SUM(COALESCE(`库存数量`, 0)) AS stock,
          SUM(COALESCE(`可用库存`, 0)) AS available_stock,
          `零售价` AS retail_price, `含税价` AS tax_price,
          MAX(`updatetime`) AS updated_at
        FROM `{source_table}` WHERE {where}
        GROUP BY {group_columns}
        ORDER BY stock DESC, product_code, warehouse
    """), params).mappings().all()
    prices = current_prices(admin)
    rows = []
    latest = None
    for row in stock_rows:
        sku = str(row["product_code"] or "").strip()
        stock = int(row["stock"] or 0)
        revision = prices.get(sku)
        values = {"retail": row["retail_price"], "tax": row["tax_price"], "core": revision.price if revision else None}
        item = {
            "product_code": sku, "product_name": row["product_name"], "brand": row["brand"],
            "warehouse": row["warehouse"], "product_type": row["product_type"],
            "stock": stock, "available_stock": int(row["available_stock"] or 0),
            "core_effective_date": revision.effective_date.isoformat() if revision else None,
        }
        for key, value in values.items():
            valid = value is not None and Decimal(str(value)) > 0
            item[f"{key}_price"] = float(value) if valid else None
            item[f"{key}_amount"] = round(stock * float(value), 2) if valid else None
        if row["updated_at"] and (latest is None or row["updated_at"] > latest):
            latest = row["updated_at"]
        rows.append(item)
    if missing_price != "all":
        rows = [row for row in rows if row[f"{missing_price}_price"] is None]
    metrics = {key: {"amount": 0.0, "priced_stock": 0, "missing_stock": 0, "priced_rows": 0, "missing_rows": 0}
               for key in ("retail", "tax", "core")}
    for row in rows:
        for key in metrics:
            valid = row[f"{key}_price"] is not None
            bucket = metrics[key]
            bucket["priced_rows" if valid else "missing_rows"] += 1
            bucket["priced_stock" if valid else "missing_stock"] += row["stock"]
            if valid:
                bucket["amount"] += row[f"{key}_amount"]
    for bucket in metrics.values():
        bucket["amount"] = round(bucket["amount"], 2)
    return ok({
        "metrics": metrics,
        "stock": sum(row["stock"] for row in rows),
        "available_stock": sum(row["available_stock"] for row in rows),
        "updated_at": latest.isoformat() if latest else None,
        "pagination": {"page": page, "page_size": page_size, "total": len(rows)},
        "rows": rows[(page - 1) * page_size:page * page_size],
    })


@router.get("/export", dependencies=[Depends(require_permission("data.export"))])
def export_valuation(
    keyword: str = "", warehouse: list[str] | None = Query(None),
    product_type: list[str] | None = Query(None), brand: list[str] | None = Query(None),
    scope: str = Query("total", pattern="^(total|split)$"),
    missing_price: str = Query("all", pattern="^(all|retail|tax|core)$"),
    ods: Session = Depends(get_ods_db), admin: Session = Depends(get_db),
) -> Response:
    data = valuation(keyword=keyword, warehouse=warehouse, product_type=product_type,
                     brand=brand, scope=scope, missing_price=missing_price, page=1, page_size=100_000,
                     ods=ods, admin=admin)["data"]
    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    sheet = workbook.add_worksheet("总库存估算" if scope == "total" else "分仓库存估算")
    headers = [
        ("product_code", "货品编号"), ("product_name", "货品名称"), ("brand", "品牌"),
        ("product_type", "分类"), ("warehouse", "仓库"), ("stock", "库存数量"),
        ("available_stock", "可用库存"), ("retail_price", "零售价"),
        ("tax_price", "含税价"), ("core_price", "核心成本价"),
        ("retail_amount", "零售价估值"), ("tax_amount", "含税价估值"),
        ("core_amount", "核心成本估值"), ("core_effective_date", "核心成本价生效日期"),
    ]
    header_format = workbook.add_format({"bold": True, "bg_color": "#E8F2EC"})
    money_format = workbook.add_format({"num_format": "#,##0.00"})
    sheet.write_row(0, 0, [label for _, label in headers], header_format)
    sheet.set_column(0, 0, 20)
    sheet.set_column(1, 1, 38)
    sheet.set_column(2, 4, 18)
    sheet.set_column(5, len(headers) - 1, 17)
    for index, row in enumerate(data["rows"], start=1):
        for column, (key, _) in enumerate(headers):
            value = row.get(key)
            if value is None:
                continue
            if key in {"product_code", "product_name", "brand", "product_type", "warehouse", "core_effective_date"}:
                sheet.write_string(index, column, str(value))
            elif key.endswith("_price") or key.endswith("_amount"):
                sheet.write_number(index, column, float(value), money_format)
            else:
                sheet.write_number(index, column, int(value))
    sheet.autofilter(0, 0, len(data["rows"]), len(headers) - 1)
    sheet.freeze_panes(1, 2)
    workbook.close()
    filename = quote(f"{'总库存' if scope == 'total' else '分仓库存'}价格估算_{date.today():%Y%m%d}.xlsx")
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})


def catalog_prices(ods, admin, keyword="", brand=None, product_type=None, source="all",
                   min_price=None, max_price=None, min_discount=None, max_discount=None):
    active, deleted = price_states(admin)
    prices = deleted if source == "delete" else active
    catalog = load_catalog(ods)
    codes = set(prices) if source == "delete" else set(catalog) | set(prices)
    by_barcode = {}
    for product in catalog.values():
        if product.get("barcode") and " / " not in product["barcode"]:
            by_barcode.setdefault(product["barcode"], []).append(product)
    selected_brands = set(brand or [])
    selected_types = set(product_type or []) - {"__all__"}
    term = keyword.strip().casefold()
    rows = []
    for code in sorted(codes):
        revision = prices.get(code)
        product = catalog.get(code)
        if revision and revision.barcode and len(by_barcode.get(revision.barcode, [])) == 1:
            product = by_barcode[revision.barcode][0]
        product = product or {"product_code": code, "product_name": revision.product_name if revision else ""}
        row = {"id": None, "price": None, "source": "", "discount": None,
               "pricing_mode": "fixed", "base_retail_price": None, **product}
        if revision:
            row.update(price_data(revision))
            row["product_name"] = product.get("product_name") or revision.product_name
            row["barcode"] = product.get("barcode") or revision.barcode
        row.setdefault("brand", "")
        row.setdefault("product_type", "")
        row.setdefault("retail_price", None)
        row.setdefault("tax_price", None)
        row["base_changed"] = bool(revision and row["pricing_mode"] == "discount" and (
            row.get("retail_price_conflict") or row["retail_price"] is None or row["base_retail_price"] is None or
            Decimal(str(row["base_retail_price"])) != Decimal(str(row["retail_price"]))
        ))
        if term and not any(term in str(row.get(key, "")).casefold() for key in ("product_code", "product_name", "barcode")):
            continue
        if selected_brands and row["brand"] not in selected_brands:
            continue
        if selected_types and not selected_types.intersection(row["product_type"].split(" / ")):
            continue
        if source == "priced" and not revision:
            continue
        if source == "missing" and revision:
            continue
        if source in {"manual", "import", "batch", "restore"} and row["source"] != source:
            continue
        if source == "changed" and not row["base_changed"]:
            continue
        if source == "conflict" and not row.get("retail_price_conflict"):
            continue
        if min_price is not None and (row["price"] is None or Decimal(str(row["price"])) < min_price):
            continue
        if max_price is not None and (row["price"] is None or Decimal(str(row["price"])) > max_price):
            continue
        row["reference_discount"] = (float((Decimal(str(row["price"])) / Decimal(str(row["retail_price"])) * 10).quantize(Decimal("0.0001")))
                                     if revision and source != "delete" and row["price"] is not None
                                     and row["retail_price"] and not row.get("retail_price_conflict") else None)
        percentage = Decimal(str(row["discount"] if row["discount"] is not None else row["reference_discount"])) * 10 if (row["discount"] is not None or row["reference_discount"] is not None) and source != "delete" else None
        if min_discount is not None and (percentage is None or percentage < min_discount):
            continue
        if max_discount is not None and (percentage is None or percentage > max_discount):
            continue
        for key in ("retail_price", "tax_price"):
            row[key] = float(row[key]) if row[key] is not None else None
        rows.append(row)
    return rows


@router.get("/prices", dependencies=[Depends(require_permission("inventory.price.manage"))])
def list_prices(
    keyword: str = "", brand: list[str] | None = Query(None), product_type: list[str] | None = Query(None),
    source: str = Query("all", pattern="^(all|priced|missing|changed|conflict|delete|manual|import|batch|restore)$"),
    min_price: Decimal | None = None, max_price: Decimal | None = None,
    min_discount: Decimal | None = Query(None, ge=0), max_discount: Decimal | None = Query(None, ge=0),
    page: int = Query(1, ge=1), page_size: int = Query(50, ge=10, le=100),
    sort_by: str = Query("product_code", pattern="^(barcode|product_code|product_name|brand|product_type|retail_price|tax_price|discount|price)$"),
    sort_order: str = Query("asc", pattern="^(asc|desc)$"),
    ods: Session = Depends(get_ods_db), admin: Session = Depends(get_db),
) -> dict:
    if min_discount is not None and max_discount is not None and min_discount > max_discount:
        raise HTTPException(422, "最低折扣不能高于最高折扣")
    rows = catalog_prices(ods, admin, keyword, brand, product_type, source, min_price, max_price, min_discount, max_discount)
    def sort_value(row):
        return (row.get("discount") if row.get("discount") is not None else row.get("reference_discount")) if sort_by == "discount" else row.get(sort_by)
    populated = [row for row in rows if sort_value(row) is not None]
    missing = [row for row in rows if sort_value(row) is None]
    populated.sort(key=sort_value, reverse=sort_order == "desc")
    rows = populated + missing
    return ok({"rows": rows[(page - 1) * page_size:page * page_size],
               "pagination": {"page": page, "page_size": page_size, "total": len(rows)}})


PRICE_COLUMNS = [("product_code", "商品编号"), ("product_name", "商品名称"), ("brand", "品牌"),
                 ("product_type", "分类"), ("retail_price", "零售价"), ("tax_price", "含税价"),
                 ("discount", "核心成本折扣"), ("price", "核心成本价")]


@router.get("/prices/export", dependencies=[Depends(require_permission("inventory.price.manage"))])
def export_prices(
    keyword: str = "", brand: list[str] | None = Query(None), product_type: list[str] | None = Query(None),
    source: str = "all", min_price: Decimal | None = None, max_price: Decimal | None = None,
    min_discount: Decimal | None = Query(None, ge=0), max_discount: Decimal | None = Query(None, ge=0),
    ods: Session = Depends(get_ods_db), admin: Session = Depends(get_db),
) -> Response:
    if min_discount is not None and max_discount is not None and min_discount > max_discount:
        raise HTTPException(422, "最低折扣不能高于最高折扣")
    rows = catalog_prices(ods, admin, keyword, brand, product_type, source, min_price, max_price, min_discount, max_discount)
    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False})
    sheet = workbook.add_worksheet("核心成本价")
    header = workbook.add_format({"bold": True, "bg_color": "#E8F2EC"})
    money = workbook.add_format({"num_format": "#,##0.00"})
    discount = workbook.add_format({"num_format": '0.####"折"'})
    sheet.write_row(0, 0, [label for _, label in PRICE_COLUMNS], header)
    sheet.set_column(0, 0, 22)
    sheet.set_column(1, 1, 42)
    sheet.set_column(2, 7, 19)
    for index, row in enumerate(rows, start=1):
        for column, (key, _) in enumerate(PRICE_COLUMNS):
            value = row.get(key)
            if key in {"price", "discount"} and row["source"] == "delete":
                value = None
            if value is not None:
                if column < 4:
                    sheet.write_string(index, column, str(value))
                else:
                    if value is not None:
                        sheet.write_number(index, column, float(value), discount if key == "discount" else money)
    sheet.autofilter(0, 0, len(rows), 7)
    sheet.freeze_panes(1, 2)
    workbook.close()
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(f"核心成本价_{date.today():%Y%m%d}.xlsx")})


@router.get("/prices/{product_code}/history", dependencies=[Depends(require_permission("inventory.price.manage"))])
def price_history(product_code: str, admin: Session = Depends(get_db)) -> dict:
    rows = admin.query(CoreCostPriceRevision).filter(CoreCostPriceRevision.product_code == product_code).order_by(CoreCostPriceRevision.id.desc()).all()
    return ok([price_data(item) for item in rows])


@router.post("/prices", dependencies=[Depends(require_permission("inventory.price.manage"))])
def save_price(item: PriceInput, user: User = Depends(get_current_user), admin: Session = Depends(get_db), ods: Session = Depends(get_ods_db)) -> dict:
    resolved = resolve_price(item, load_catalog(ods, [item.product_code]), current_prices(admin))
    add_revision(admin, resolved, user.username, "manual")
    admin.commit()
    return ok({"saved": 1})


def revision_input(revision):
    return PriceInput(product_code=revision.product_code, product_name=revision.product_name, barcode=revision.barcode,
                      price=revision.price, effective_date=date.today(),
                      pricing_mode=revision.pricing_mode or "fixed", discount=revision.discount,
                      base_retail_price=revision.base_retail_price)


@router.delete("/prices/{product_code}", dependencies=[Depends(require_permission("inventory.price.manage"))])
def delete_price(product_code: str, user: User = Depends(get_current_user), admin: Session = Depends(get_db)) -> dict:
    current = current_prices(admin).get(product_code)
    if current is None:
        raise HTTPException(404, "当前核心成本价不存在")
    add_revision(admin, revision_input(current), user.username, "delete")
    admin.commit()
    return ok({"deleted": 1})


@router.post("/prices/{product_code}/restore", dependencies=[Depends(require_permission("inventory.price.manage"))])
def restore_price(product_code: str, revision_id: int, user: User = Depends(get_current_user), admin: Session = Depends(get_db)) -> dict:
    revision = admin.get(CoreCostPriceRevision, revision_id)
    if revision is None or revision.product_code != product_code or revision.source == "delete":
        raise HTTPException(404, "价格版本不存在")
    add_revision(admin, revision_input(revision), user.username, "restore", revision.id)
    admin.commit()
    return ok({"restored": 1})


@router.post("/prices/import-preview", dependencies=[Depends(require_permission("inventory.price.manage"))])
async def import_preview(file: UploadFile = File(...), ods: Session = Depends(get_ods_db), admin: Session = Depends(get_db)) -> dict:
    content = await file.read(5_000_001)
    if len(content) > 5_000_000 or not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "请上传不超过 5 MB 的 xlsx 文件")
    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        iterator = workbook.active.values
        header = next(iterator)
        columns = {str(value).strip(): index for index, value in enumerate(header) if value is not None}
        code_key = "商品编号" if "商品编号" in columns else "货品编号"
        name_key = "商品名称" if "商品名称" in columns else "货品名称"
        if code_key not in columns or ("核心成本价" not in columns and "核心成本折扣" not in columns):
            raise HTTPException(400, "请提供商品编号及核心成本价或核心成本折扣列，兼容旧版货品编号")
        raw_rows = []
        for number, raw in enumerate(iterator, 2):
            if number > 5001:
                raise HTTPException(400, "一次最多导入 5000 行")
            if any(value is not None for value in raw):
                raw_rows.append((number, raw))
        def value(raw, key):
            index = columns.get(key)
            return raw[index] if index is not None and index < len(raw) else None
        codes = [str(value(raw, code_key) or "").strip() for _, raw in raw_rows]
        catalog = load_catalog(ods, list(set(codes)))
        existing = current_prices(admin)
        rows, errors, seen = [], [], set()
        blank = 0
        for (number, raw), code in zip(raw_rows, codes):
            if not code or code in seen or len(code) > 128:
                errors.append(f"第 {number} 行商品编号为空、重复或过长")
                continue
            seen.add(code)
            price, discount = value(raw, "核心成本价"), value(raw, "核心成本折扣")
            if price in (None, "") and discount in (None, ""):
                blank += 1
                continue
            try:
                # Imported explicit amounts are authoritative, including exports
                # whose reference discount or retail basis differs locally.
                discount = (Decimal(str(discount).strip().removesuffix("折"))
                            if price in (None, "") and discount not in (None, "") else None)
                base = value(raw, "零售价") if discount is not None else None
                item = PriceInput(product_code=code, product_name=str(value(raw, name_key) or "")[:255],
                                  price=price if price not in (None, "") else None,
                                  pricing_mode="discount" if discount is not None else "fixed", discount=discount,
                                  base_retail_price=base if base not in (None, "") else None,
                                  expected_revision_id=existing[code].id if code in existing else None)
                resolved = resolve_price(item, catalog, existing)
                rows.append({**resolved.model_dump(mode="json"), "brand": catalog.get(code, {}).get("brand", str(value(raw, "品牌") or ""))})
            except HTTPException as exc:
                errors.append(f"第 {number} 行：{exc.detail}")
            except (InvalidOperation, ValueError):
                errors.append(f"第 {number} 行核心成本价或折扣无效")
        return ok({"rows": rows, "errors": errors, "blank_price_count": blank})
    except HTTPException:
        raise
    except (ValueError, KeyError, StopIteration, OSError, BadZipFile) as exc:
        raise HTTPException(400, "Excel 格式或表头不正确") from exc
    finally:
        if workbook:
            workbook.close()


def persist_prices(payload, user, admin, ods, source):
    codes = [item.product_code for item in payload.rows]
    if len(codes) != len(set(codes)):
        raise HTTPException(400, "导入数据包含重复商品编号")
    catalog, existing = load_catalog(ods, codes), current_prices(admin)
    # Validate the entire selection before writing any revisions.
    resolved = [resolve_price(item, catalog, existing) for item in payload.rows]
    saved = 0
    for item in resolved:
        previous = existing.get(item.product_code)
        if previous and all(getattr(previous, key) == getattr(item, key) for key in ("price", "pricing_mode", "discount", "base_retail_price")):
            continue
        add_revision(admin, item, user.username, source)
        saved += 1
    admin.commit()
    return ok({"saved": saved, "skipped": len(codes) - saved})


@router.post("/prices/import", dependencies=[Depends(require_permission("inventory.price.manage"))])
def import_prices(payload: ImportInput, user: User = Depends(get_current_user), admin: Session = Depends(get_db), ods: Session = Depends(get_ods_db)) -> dict:
    return persist_prices(payload, user, admin, ods, "import")


@router.post("/prices/batch-preview", dependencies=[Depends(require_permission("inventory.price.manage"))])
def batch_preview(payload: BatchPriceInput, ods: Session = Depends(get_ods_db), admin: Session = Depends(get_db)) -> dict:
    if not payload.brand and not payload.product_type and not payload.keyword.strip():
        raise HTTPException(422, "批量设置前请先选择品牌、分类或输入关键词，避免影响全部商品")
    existing = current_prices(admin)
    rows = catalog_prices(ods, admin, payload.keyword, payload.brand, payload.product_type)
    if len(rows) > 5000:
        raise HTTPException(422, "匹配超过 5000 个商品，请缩小筛选范围")
    preview, errors = [], []
    skipped = 0
    for row in rows:
        if row["price"] is not None and row["pricing_mode"] == "fixed" and not payload.overwrite_fixed:
            skipped += 1
            continue
        if row.get("retail_price_conflict") or row["retail_price"] is None:
            errors.append(row["product_code"] + ("：零售价冲突" if row.get("retail_price_conflict") else "：缺少零售价"))
            continue
        try:
            price = calculate_cost(row["retail_price"], payload.discount)
        except HTTPException as exc:
            errors.append(f"{row['product_code']}：{exc.detail}")
            continue
        item = PriceInput(product_code=row["product_code"], product_name=row["product_name"][:255],
                          pricing_mode="discount", discount=payload.discount,
                          price=price,
                          base_retail_price=row["retail_price"],
                          expected_revision_id=existing[row["product_code"]].id if row["product_code"] in existing else None)
        preview.append({**item.model_dump(mode="json"), "old_price": row["price"], "brand": row["brand"]})
    return ok({"rows": preview, "errors": errors, "skipped": skipped, "matched": len(rows)})


@router.post("/prices/batch", dependencies=[Depends(require_permission("inventory.price.manage"))])
def batch_save(payload: ImportInput, user: User = Depends(get_current_user), admin: Session = Depends(get_db), ods: Session = Depends(get_ods_db)) -> dict:
    return persist_prices(payload, user, admin, ods, "batch")
